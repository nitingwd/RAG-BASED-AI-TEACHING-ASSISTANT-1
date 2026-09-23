import os
import re
import json
import io
import tempfile
import subprocess
import shutil
import time
from urllib.parse import urlparse
from typing import Optional, List

import streamlit as st
from dotenv import load_dotenv
from groq import Groq
from gtts import gTTS
from PIL import Image
from docx import Document
from docx.shared import Inches

from langchain_groq import ChatGroq
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, CSVLoader, TextLoader

load_dotenv()

# ============================================================
# CONFIG
# ============================================================
MODEL_NAME = os.getenv("GROQ_MODEL", "openai/gpt-oss-20b")
EMBEDDING_MODEL = os.getenv(
    "EMBEDDING_MODEL", "sentence-transformers/all-MiniLM-L6-v2"
)
MAX_CONTEXT_CHARS = 12000


def get_api_key():
    """Read Groq API key from Streamlit secrets first, then environment."""
    try:
        key = st.secrets.get("GROQ_API_KEY")
        if key:
            return key
    except Exception:
        pass
    return os.getenv("GROQ_API_KEY")


def _require_api_key():
    key = get_api_key()
    if not key:
        raise RuntimeError(
            "GROQ_API_KEY missing. Add it to Streamlit Secrets or your .env file."
        )
    return key


@st.cache_resource(show_spinner=False)
def get_embeddings():
    return HuggingFaceEmbeddings(
        model_name=EMBEDDING_MODEL,
        model_kwargs={"device": "cpu"},
    )


def get_lang_code(lang_name):
    mapping = {
        "Hindi": "hi",
        "Gujarati": "gu",
        "Hinglish": "hi",
        "English": "en",
        "Marathi": "mr",
    }
    return mapping.get(lang_name, "hi")


def _llm(temperature=0.2, max_tokens=None):
    kwargs = {
        "model": MODEL_NAME,
        "groq_api_key": _require_api_key(),
        "temperature": temperature,
    }
    if max_tokens is not None:
        kwargs["max_tokens"] = int(max_tokens)
    try:
        return ChatGroq(**kwargs)
    except TypeError:
        # Older langchain-groq versions may not expose max_tokens in the
        # constructor. Fall back to the normal client instead of crashing.
        kwargs.pop("max_tokens", None)
        return ChatGroq(**kwargs)


def _clean_text(text):
    text = re.sub(r"<[^>]+>", " ", str(text))
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _invoke(prompt, temperature=0.2, max_tokens=None):
    return _llm(temperature, max_tokens=max_tokens).invoke(prompt).content


def _invoke_long(prompt, temperature=0.25, max_tokens=7000):
    """Long-form generation helper used for complete lessons and project code."""
    try:
        return str(_invoke(prompt, temperature=temperature, max_tokens=max_tokens) or "").strip()
    except Exception:
        # Compatibility fallback for older client/model combinations.
        return str(_invoke(prompt, temperature=temperature) or "").strip()


def _json_from_text(text, default=None):
    """Extract the first JSON object/array from an LLM response."""
    if not text:
        return default
    text = text.strip()
    candidates = []

    first_array = text.find("[")
    last_array = text.rfind("]")
    if first_array >= 0 and last_array > first_array:
        candidates.append(text[first_array:last_array + 1])

    first_obj = text.find("{")
    last_obj = text.rfind("}")
    if first_obj >= 0 and last_obj > first_obj:
        candidates.append(text[first_obj:last_obj + 1])

    for candidate in candidates:
        try:
            return json.loads(candidate)
        except Exception:
            continue
    return default


# ============================================================
# AUDIO
# ============================================================

def text_to_audio_file(text, language_name="Hinglish"):
    try:
        clean = _clean_text(text)[:3500]
        if not clean:
            return None

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".mp3")
        tmp.close()
        try:
            gTTS(text=clean, lang=get_lang_code(language_name), slow=False).save(tmp.name)
        except Exception:
            # gTTS can reject some language combinations; English is a safe fallback.
            gTTS(text=clean, lang="en", slow=False).save(tmp.name)

        if os.path.exists(tmp.name) and os.path.getsize(tmp.name) > 1000:
            return tmp.name
        return None
    except Exception as e:
        print(f"TTS Error: {e}")
        return None


# ============================================================
# AI VIDEO LESSON GENERATOR
# ============================================================

def _find_font(language_name="English", bold=False):
    """Find a Unicode font that works on Streamlit Cloud/Linux."""
    candidates = []
    lang = str(language_name or "English").lower()
    if "gujar" in lang:
        candidates += ["/usr/share/fonts/truetype/noto/NotoSansGujarati-Regular.ttf",
                       "/usr/share/fonts/opentype/noto/NotoSansGujarati-Regular.ttf"]
    elif any(x in lang for x in ["hindi", "marathi"]):
        candidates += ["/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",
                       "/usr/share/fonts/opentype/noto/NotoSansDevanagari-Regular.ttf"]
    candidates += [
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            return path
    return None


def _wrap_for_slide(text, font, max_width, draw):
    words = str(text or "").split()
    lines, current = [], ""
    for word in words:
        trial = word if not current else current + " " + word
        try:
            width = draw.textbbox((0, 0), trial, font=font)[2]
        except Exception:
            width = len(trial) * 16
        if width <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def _make_video_slide(title, bullets, scene_no, total, language_name, out_path):
    """Create a clean 16:9 educational slide using Pillow."""
    from PIL import ImageDraw, ImageFont

    W, H = 1280, 720
    img = Image.new("RGB", (W, H), (15, 23, 42))
    draw = ImageDraw.Draw(img)
    regular_path = _find_font(language_name)
    if not regular_path:
        raise RuntimeError("Unicode font not available for video slides")

    try:
        title_font = ImageFont.truetype(regular_path, 42)
        body_font = ImageFont.truetype(regular_path, 27)
        small_font = ImageFont.truetype(regular_path, 20)
    except Exception:
        title_font = ImageFont.load_default()
        body_font = ImageFont.load_default()
        small_font = ImageFont.load_default()

    # Simple classroom-style cards; no external images are required.
    draw.rounded_rectangle((55, 45, W - 55, H - 45), radius=28, fill=(30, 41, 59), outline=(71, 85, 105), width=2)
    draw.text((85, 72), "EduSolve AI  •  AI Video Lesson", font=small_font, fill=(148, 163, 184))
    draw.text((85, 112), str(title)[:70], font=title_font, fill=(248, 250, 252))

    y = 205
    for bullet in bullets[:6]:
        lines = _wrap_for_slide(bullet, body_font, 1030, draw)
        if not lines:
            continue
        draw.ellipse((92, y + 10, 106, y + 24), fill=(56, 189, 248))
        draw.text((125, y), lines[0], font=body_font, fill=(226, 232, 240))
        y += 48
        for line in lines[1:3]:
            draw.text((125, y), line, font=body_font, fill=(203, 213, 225))
            y += 40
        y += 10
        if y > 610:
            break

    progress = max(1, min(total, scene_no)) / max(1, total)
    draw.rounded_rectangle((85, 665, 1195, 676), radius=6, fill=(51, 65, 85))
    draw.rounded_rectangle((85, 665, 85 + int(1110 * progress), 676), radius=6, fill=(56, 189, 248))
    draw.text((1100, 620), f"{scene_no}/{total}", font=small_font, fill=(148, 163, 184))
    img.save(out_path, "PNG")


def _video_script_fallback(topic, language):
    """Robust offline-safe lesson plan that still teaches a topic end-to-end."""
    return {
        "title": f"{topic} — Complete Visual Lesson",
        "scenes": [
            {"heading":"1. Hook & Learning Goal","type":"hook","narration":f"Aaj hum {topic} ko sirf ek definition ke roop mein nahi, balki ek complete concept ke roop mein samjhenge. Sabse pehle dekhenge ki problem kya hai aur ye topic important kyun hai. Phir exact definition aur intuition, main parts, working, ek detailed example, practical use, common mistakes aur final recap cover karenge. Lesson ka goal hai ki aap video ke end mein {topic} ko apni language mein explain kar saken, uska working sequence bata saken aur ek new question par apply kar saken.","points":["Why this topic matters","Problem being solved","Learning goal"],"visual_type":"concept"},
            {"heading":"2. Exact Definition & Intuition","type":"concept","narration":f"Ab {topic} ki exact definition samjho. Pehle simple language mein meaning samjhenge aur phir technical wording se connect karenge. Kisi bhi concept ko samajhne ke liye teen cheezein identify karo: input kya hai, process ya rule kya hai, aur output kya milta hai. Is mental model se difficult terminology bhi easy ho jati hai. Definition ko real-life intuition ke saath connect karo, kyunki sirf definition yaad karne ke bajay uska meaning samajhna long-term learning ke liye important hai.","points":["Exact meaning","Simple intuition","Input → process → output"],"visual_type":"concept"},
            {"heading":"3. Main Components","type":"architecture","narration":f"Ab {topic} ko uske essential components mein break karte hain. Har component ka naam, role aur doosre components ke saath relationship samjho. Ek component ko isolated fact ki tarah yaad mat karo; dekho information ya control kis direction mein move hota hai. Is visual map se aapko topic ka structure ek mental picture ki tarah yaad rahega. Agar kisi component ko remove kar diya jaye to system ya method par kya effect padega, is question se bhi understanding check kar sakte hain.","points":["Components","Role of each part","Relationships","Information flow"],"visual_type":"architecture"},
            {"heading":"4. How It Works — Step by Step","type":"process","narration":f"Ab {topic} ka actual working sequence follow karo. Starting input se begin karo, phir har operation, decision, transformation aur output ko order mein dekho. Har step ke saath do questions poochho: ye step kya karta hai aur ye step kyun zaroori hai? Isse ratta lagane ke bajay cause-and-effect samajh aayega. Agar process mein conditions, branches, iterations ya stages hain, to unhe clearly distinguish karo. End mein poora flow beginning se output tak ek baar mentally repeat karo.","points":["Input","Operation","Decision / transformation","Output"],"visual_type":"process"},
            {"heading":"5. Detailed Worked Example","type":"example","narration":f"Ab {topic} ko ek concrete worked example par apply karte hain. Pehle given information identify karo, phir relevant rule, method ya component choose karo. Uske baad har intermediate step ko explain karte hue result tak pahunchte hain. Sirf final answer dekhna enough nahi hai; important ye hai ki har step previous step se logically kaise connected hai. Isi example ko thoda change karke socho aur dekho ki method ka behavior kaise badalta hai. Ye transfer of learning exam aur practical problem solving dono mein useful hai.","points":["Given","Method","Intermediate steps","Final result"],"visual_type":"example"},
            {"heading":"6. Visual / Formula / Code View","type":"practical","narration":f"Ab topic ko us representation mein dekho jo uske liye sabse useful hai. Programming topic mein input, logic, important code blocks aur output ka relation samjho. Mathematics ya electronics mein symbols, equations, substitutions aur result ko connect karo. Architecture-based topic mein blocks aur data flow dekho. Is scene ka purpose theory ko visible banana hai, taaki aap abstract explanation ko actual implementation, calculation ya system behavior se connect kar saken.","points":["Representation","Logic / formula","Implementation","Result"],"visual_type":"practical"},
            {"heading":"7. Second Example / Edge Case","type":"example","narration":f"Ek second example ya edge case se {topic} ki understanding ko test karte hain. Is baar situation mein ek important change hoga. Pehle predict karo ki output ya behavior kya hoga, phir step by step verify karo. Edge cases samajhne se pata chalta hai ki concept sirf ideal example ke liye nahi hai. Agar result expected se different ho, to identify karo ki kaunsa assumption change hua. Isi habit se conceptual clarity aur problem-solving accuracy improve hoti hai.","points":["New situation","Prediction","Verification","Edge case"],"visual_type":"example"},
            {"heading":"8. Comparison & Common Confusion","type":"comparison","narration":f"Students ko {topic} samajhte waqt similar terms, methods ya components ke beech confusion ho sakta hai. Isliye ab relevant comparison dekho: purpose, working, input, output, use-case aur limitation ke basis par difference identify karo. Sirf terminology ka difference yaad mat karo. Practical situation mein kaunsa option kab use hoga aur kyun, ye samajhna zyada important hai. Common misconception ko bhi identify karo aur correct mental model ke saath replace karo.","points":["Similar concepts","Key differences","When to use","Common mistake"],"visual_type":"comparison"},
            {"heading":"9. Applications & Limitations","type":"practical","narration":f"Ab {topic} ke real-world applications dekho. Concept ka use kahan hota hai, kis type ki problem solve karta hai aur practical system mein iska role kya hota hai, ye samjho. Saath hi limitations ko ignore mat karo. Har method ya technology ke assumptions, constraints, cost, performance, accuracy ya hardware/software dependencies ho sakti hain. Application ko limitation ke saath samajhne se aapko pata chalega ki concept kab appropriate hai aur kab alternative approach consider karni chahiye.","points":["Applications","Benefits","Limitations","Selection criteria"],"visual_type":"concept"},
            {"heading":"10. Exam / Viva Focus","type":"recap","narration":f"Exam aur viva ke point of view se {topic} ko compress karte hain. Aapko minimum ye explain kar pana chahiye: definition kya hai, problem kya solve karta hai, main components kaun se hain, working ka sequence kya hai, example kaise solve hota hai, applications kya hain aur limitations kya hain. Agar diagram, algorithm, formula ya code relevant hai to uska purpose bhi explain karo. Kisi bhi short-answer question ke liye definition plus key working, aur long-answer ke liye diagram plus detailed example ka structure useful rahega.","points":["Definition","Diagram / working","Example","Applications","Limitations"],"visual_type":"recap"},
            {"heading":"11. Complete Mental Map","type":"architecture","narration":f"Ab poore {topic} ko ek single mental map mein connect karo. Problem se start karo, definition tak jao, definition ko components se connect karo, components ko working flow se, working ko example se aur example ko application se. Phir limitation aur common confusion ko attach karo. Jab ye chain clear ho jaye, to topic disconnected facts ka collection nahi rehta; ek connected system ban jata hai. Video ke baad bina notes dekhe isi chain ko khud bolkar reproduce karna best self-test hai.","points":["Problem → concept","Components → working","Example → application","Limitation → decision"],"visual_type":"architecture"},
            {"heading":"12. Final Recap & Self Check","type":"recap","narration":f"Final recap mein {topic} ko paanch questions se test karo. Pehla: ye kya hai aur exact definition kya hai? Doosra: iski need kyun hai? Teesra: ye step by step kaise work karta hai? Chautha: ek practical example mein ise kaise apply karoge? Paanchva: common mistake, limitation ya edge case kya hai? Agar aap in paanch questions ka answer bina video dekhe de sakte hain, to topic ki core understanding strong hai. Agar kisi answer mein gap hai, wahi aapka revision point hai.","points":["What?","Why?","How?","Example?","Limitation / mistake?"],"visual_type":"recap"}
        ],
    }

def generate_ai_video_script(topic, language="Hinglish", style="Animated Classroom"):
    """Generate a deep, visual-first teaching script with scene instructions."""
    topic = _clean_text(topic)
    if not topic:
        raise ValueError("Topic required")
    if re.search(r"\besp32\b", topic, re.I):
        topic_lock = """ESP32 TOPIC LOCK: The lesson must be ONLY about ESP32. Cover what ESP32 is; family/board context without inventing board-specific specifications; high-level architecture; GPIO; ADC/DAC where applicable; PWM; UART/I2C/SPI; Wi-Fi and Bluetooth/BLE; memory at a high level; power considerations; programming/toolchain; a simple ESP32 program; an IoT sensor example; applications; limitations/safety; and recap. Explicitly say when a feature depends on the exact ESP32 variant/board. Never replace ESP32 with Arduino UNO, Raspberry Pi, NodeMCU/ESP8266, or another controller."""
    else:
        topic_lock = f"TOPIC LOCK: Every scene must directly teach '{topic}'. Do not drift into a related product, framework, or neighboring topic unless it is explicitly introduced only as a comparison."

    prompt = f"""You are EduSolve AI's senior educational video director and expert teacher.
Create a COMPLETE, self-contained educational video lesson for EXACTLY this topic: {topic}
{topic_lock}
Language: {language}
Visual style: {style}

The student has specifically complained that short AI videos mention a topic but do not actually explain it. Therefore this lesson MUST teach the requested topic completely enough for a college student to understand it without needing a second video. Do not produce a generic template with the topic name inserted.

FIRST classify the topic internally (algorithm, programming, AI/ML, DBMS, networking, OS, compiler, electronics/IoT, mathematics, science, business, etc.) and then cover the canonical concepts normally required to understand THAT exact topic. Do not add unrelated syllabus topics.

Create 10-14 coherent scenes. The sequence should normally be:
1) hook/problem and learning goal,
2) exact definition,
3) intuition/analogy,
4) components or prerequisites,
5) core working step-by-step,
6) diagram/architecture or formula/code representation when relevant,
7) detailed worked example,
8) second example or edge case,
9) comparison/common misconceptions,
10) applications and limitations,
11) exam/viva/practical takeaways,
12) connected recap + self-check; add extra scenes whenever the topic genuinely requires them.

For each scene return:
- heading: short but specific title
- type: hook|concept|architecture|process|example|comparison|practical|code|formula|recap
- narration: natural spoken teaching, roughly 150-230 words; explain WHAT, WHY and HOW, not just bullet points
- points: 3-6 short reinforcement labels
- visual_type: concept|process|architecture|comparison|example|code|formula|timeline|recap
- visual_elements: 3-8 concrete things that should visibly appear
- presenter_action: one short action such as pointing, writing, comparing, tracing a flow, or demonstrating

NON-NEGOTIABLE TEACHING RULES:
- Never merely list facts. Explain the relationship between ideas.
- Define every important technical term before relying on it.
- If the topic has an algorithm, show every major step, a worked input, intermediate states, output, and time/space complexity when applicable.
- If it has architecture, show every important block and the direction/purpose of data flow.
- If it has a formula, define symbols, show substitution and interpretation of the result.
- If it has programming, show complete relevant code (no '...', 'rest of code', or placeholders) and explain the important lines/blocks.
- If it has electronics/IoT, explain components, connections, signals/data flow and practical operation without inventing board-specific specifications.
- If it has comparisons, explain meaningful criteria rather than superficial wording differences.
- Include at least TWO concrete examples whenever the topic permits.
- Include common mistakes, edge cases, limitations, applications, and when the concept should or should not be used when applicable.
- Use analogies only when they clarify the concept; always connect the analogy back to the technical explanation.
- Keep the topic tightly locked. Do not silently substitute a related topic.
- End with a concise connected recap and 5 self-check questions with their purpose clear from the narration.
- Never invent facts, specifications, citations, datasets, or sources.

The final video should feel like a patient expert teacher teaching the student from zero to confident understanding, not like an AI reading slides.

Return ONLY valid JSON:
{{"title":"...","scenes":[{{"heading":"...","type":"concept","narration":"...","points":["..."],"visual_type":"process","visual_elements":["..."],"presenter_action":"..."}}]}}

Return ONLY valid JSON:
{{"title":"...","scenes":[{{"heading":"...","type":"concept","narration":"...","points":["..."],"visual_type":"process","visual_elements":["..."],"presenter_action":"..."}}]}}
"""
    try:
        raw = _invoke_long(prompt, temperature=0.22, max_tokens=16000)
        data = _json_from_text(raw, default=None)
        if isinstance(data, dict) and isinstance(data.get("scenes"), list) and data.get("scenes"):
            # Normalize scene fields so malformed model output cannot break rendering.
            normalized = []
            allowed = {"concept", "process", "architecture", "comparison", "example", "code", "formula", "timeline", "recap"}
            for i, scene in enumerate(data["scenes"][:14], 1):
                if not isinstance(scene, dict):
                    continue
                points = scene.get("points", [])
                if isinstance(points, str): points = [points]
                visual_elements = scene.get("visual_elements", [])
                if isinstance(visual_elements, str): visual_elements = [visual_elements]
                normalized.append({
                    "heading": _clean_text(scene.get("heading") or f"Scene {i}"),
                    "type": _clean_text(scene.get("type") or "concept").lower(),
                    "narration": str(scene.get("narration") or "").strip(),
                    "points": [_clean_text(x) for x in points if _clean_text(x)][:6],
                    "visual_type": _clean_text(scene.get("visual_type") or "concept").lower() if _clean_text(scene.get("visual_type") or "concept").lower() in allowed else "concept",
                    "visual_elements": [_clean_text(x) for x in visual_elements if _clean_text(x)][:8],
                    "presenter_action": _clean_text(scene.get("presenter_action") or "Explain the highlighted visual.")[:180],
                })
            if normalized:
                data["scenes"] = normalized
                return data
    except Exception as e:
        print(f"Video script error: {e}")
    return _video_script_fallback(topic, language)


def _draw_teacher(draw, x=55, y=120, scale=1.0, talking=False, language_name="English"):
    """Draw a clean presenter-style human illustration for the offline fallback video."""
    from PIL import ImageDraw
    s = float(scale)
    # soft presenter panel
    draw.rounded_rectangle((int(x-20*s), int(y-20*s), int(x+275*s), int(y+535*s)), radius=int(28*s), fill=(15, 23, 42), outline=(71, 85, 105), width=max(2, int(2*s)))
    cx = x + 125*s
    head_y = y + 120*s
    # hair / head
    draw.ellipse((int(cx-60*s), int(head_y-60*s), int(cx+60*s), int(head_y+60*s)), fill=(245, 190, 150), outline=(120, 80, 60), width=max(1, int(2*s)))
    draw.pieslice((int(cx-62*s), int(head_y-65*s), int(cx+62*s), int(head_y+40*s)), 180, 360, fill=(45, 35, 30))
    # eyes
    draw.ellipse((int(cx-28*s), int(head_y-8*s), int(cx-18*s), int(head_y+2*s)), fill=(25,25,25))
    draw.ellipse((int(cx+18*s), int(head_y-8*s), int(cx+28*s), int(head_y+2*s)), fill=(25,25,25))
    # mouth changes between frames to simulate speech
    if talking:
        draw.ellipse((int(cx-14*s), int(head_y+20*s), int(cx+14*s), int(head_y+42*s)), fill=(80, 35, 40))
    else:
        draw.arc((int(cx-18*s), int(head_y+18*s), int(cx+18*s), int(head_y+42*s)), 10, 170, fill=(80, 35, 40), width=max(1, int(3*s)))
    # neck + body
    draw.rectangle((int(cx-20*s), int(head_y+58*s), int(cx+20*s), int(head_y+105*s)), fill=(235, 170, 135))
    draw.rounded_rectangle((int(cx-88*s), int(head_y+92*s), int(cx+88*s), int(head_y+305*s)), radius=int(28*s), fill=(37, 99, 235))
    # arms, one pointing toward visual area
    draw.line((int(cx+70*s), int(head_y+145*s), int(cx+155*s), int(head_y+105*s)), fill=(245, 190, 150), width=max(8, int(16*s)))
    draw.line((int(cx-70*s), int(head_y+145*s), int(cx-125*s), int(head_y+230*s)), fill=(245, 190, 150), width=max(8, int(16*s)))
    # talking label
    font_path = _find_font(language_name)
    if font_path:
        from PIL import ImageFont
        try:
            f = ImageFont.truetype(font_path, max(14, int(18*s)))
            draw.text((int(x+35*s), int(y+350*s)), "AI TEACHER", font=f, fill=(226,232,240))
            draw.text((int(x+35*s), int(y+385*s)), "explaining…", font=f, fill=(148,163,184))
        except Exception:
            pass


def _draw_scene_visual(img, scene, progress, frame_no, total_frames, language_name):
    """Render an animated teaching canvas: presenter + diagrams/flows/code/formulas."""
    from PIL import ImageDraw, ImageFont
    draw = ImageDraw.Draw(img)
    W, H = img.size
    font_path = _find_font(language_name)
    try:
        title_font = ImageFont.truetype(font_path, 36)
        body_font = ImageFont.truetype(font_path, 23)
        small_font = ImageFont.truetype(font_path, 17)
    except Exception:
        title_font = body_font = small_font = ImageFont.load_default()

    heading = _clean_text(scene.get("heading") or "Lesson")
    vtype = _clean_text(scene.get("visual_type") or scene.get("type") or "concept").lower()
    points = scene.get("points", []) if isinstance(scene.get("points", []), list) else [scene.get("points", "")]
    elements = scene.get("visual_elements", []) if isinstance(scene.get("visual_elements", []), list) else [scene.get("visual_elements", "")]
    points = [_clean_text(x) for x in points if _clean_text(x)][:5]
    elements = [_clean_text(x) for x in elements if _clean_text(x)][:6]

    img.paste((8, 15, 30), (0, 0, W, H))
    draw.rounded_rectangle((35, 28, W-35, H-28), radius=30, fill=(15, 23, 42), outline=(51,65,85), width=2)
    draw.text((55, 48), "EduSolve AI  •  Visual Masterclass", font=small_font, fill=(148,163,184))
    draw.text((55, 82), heading[:68], font=title_font, fill=(248,250,252))
    # presenter panel
    _draw_teacher(draw, 55, 145, 0.72, talking=(frame_no % 2 == 0), language_name=language_name)

    # visual board
    bx1, by1, bx2, by2 = 350, 145, W-65, H-70
    draw.rounded_rectangle((bx1, by1, bx2, by2), radius=24, fill=(30,41,59), outline=(71,85,105), width=2)
    label = {"process":"PROCESS", "architecture":"SYSTEM MAP", "comparison":"COMPARE", "example":"WORKED EXAMPLE", "code":"CODE LOGIC", "formula":"FORMULA", "timeline":"TIMELINE", "recap":"RECAP"}.get(vtype, "CONCEPT")
    draw.text((bx1+24, by1+18), label, font=small_font, fill=(56,189,248))

    # Reveal more visual content as narration progresses.
    reveal = max(1, int(round((len(elements) or len(points) or 4) * max(0.15, progress))))
    items = (elements or points or ["Main idea", "Working", "Example", "Result"])[:6]
    shown = items[:reveal]

    def box(x, y, w, h, text, fill=(51,65,85), accent=False):
        draw.rounded_rectangle((x, y, x+w, y+h), radius=16, fill=fill, outline=(56,189,248) if accent else (100,116,139), width=2)
        lines = _wrap_for_slide(text, body_font, w-28, draw)[:3]
        yy = y + 14
        for line in lines:
            draw.text((x+14, yy), line, font=body_font, fill=(241,245,249))
            yy += 30

    if vtype in ("process", "timeline"):
        y = by1 + 120
        count = len(shown)
        for i, item in enumerate(shown):
            x = bx1 + 25 + i * max(1, int((bx2-bx1-70)/max(1,count)))
            box(x, y, 120, 105, item, accent=(i == count-1))
            if i < count-1:
                x2 = x + 122
                draw.line((x2, y+52, x2+32, y+52), fill=(56,189,248), width=5)
                draw.polygon([(x2+32,y+52),(x2+20,y+44),(x2+20,y+60)], fill=(56,189,248))
    elif vtype == "comparison":
        mid = (bx1+bx2)//2
        box(bx1+25, by1+120, (bx2-bx1)//2-40, 245, "SIDE A\n" + "\n".join(points[:3]), fill=(38,50,70), accent=True)
        box(mid+15, by1+120, (bx2-bx1)//2-40, 245, "SIDE B\n" + "\n".join(points[3:6] or points[:3]), fill=(38,50,70), accent=False)
        draw.text((mid-32, by1+390), "VS", font=title_font, fill=(248,250,252))
    elif vtype == "code":
        code = "\n".join(points or elements or ["Input → logic → output"])
        draw.rounded_rectangle((bx1+25, by1+105, bx2-25, by2-35), radius=16, fill=(2,6,23), outline=(71,85,105), width=2)
        lines = code.splitlines()[:10]
        yy = by1+130
        for i,line in enumerate(lines):
            draw.text((bx1+45, yy), f"{i+1:02d}  {line[:62]}", font=small_font, fill=(226,232,240))
            yy += 30
    elif vtype == "formula":
        formula = elements[0] if elements else "Formula → substitution → result"
        draw.text((bx1+55, by1+125), formula[:55], font=title_font, fill=(248,250,252))
        y=by1+220
        for i,item in enumerate(shown):
            box(bx1+45, y+i*75, bx2-bx1-90, 55, item, accent=(i==len(shown)-1))
    else:
        # concept / architecture / example / recap: connected mental model
        center = items[0] if items else "Main Idea"
        cx=(bx1+bx2)//2; cy=by1+230
        box(cx-110, cy-50, 220, 100, center, accent=True)
        coords=[(bx1+75,by1+115),(bx2-225,by1+115),(bx1+75,by2-150),(bx2-225,by2-150)]
        for i,item in enumerate(shown[1:5]):
            x,y=coords[i]
            box(x,y,150,90,item)
            draw.line((cx,cy,x+75,y+45), fill=(100,116,139), width=3)

    # progress + speech indicator
    p = max(0.0, min(1.0, progress))
    draw.rounded_rectangle((55, H-52, W-55, H-40), radius=5, fill=(51,65,85))
    draw.rounded_rectangle((55, H-52, 55+int((W-110)*p), H-40), radius=5, fill=(56,189,248))
    # animated sound bars
    base_x = W-220
    for j in range(8):
        h = 10 + ((frame_no + j*3) % 7)*4 if frame_no % 2 == 0 else 12 + ((frame_no + j) % 4)*3
        draw.rounded_rectangle((base_x+j*18, H-92-h, base_x+10+j*18, H-92), radius=4, fill=(56,189,248))


def _silent_audio(audio_path, duration=4.0):
    subprocess.run(["ffmpeg","-y","-f","lavfi","-i",f"anullsrc=channel_layout=mono:sample_rate=24000","-t",f"{duration:.2f}","-c:a","aac","-b:a","96k",audio_path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def create_ai_video_lesson(topic, language="Hinglish", style="Animated Classroom"):
    """Create a visual-first narrated MP4 without paid avatar credits.

    The fallback is deliberately NOT a slideshow: it uses an illustrated teacher,
    speech indicator, progressive diagrams, process arrows, architecture maps,
    comparison panels, code/formula boards and scene-by-scene narration.
    """
    topic = _clean_text(topic)
    if not topic:
        raise ValueError("Topic required")
    if not shutil.which("ffmpeg") or not shutil.which("ffprobe"):
        raise RuntimeError("FFmpeg/FFprobe is not available. Add 'ffmpeg' to packages.txt and redeploy.")

    script = generate_ai_video_script(topic, language, style)
    scenes = script.get("scenes", []) if isinstance(script, dict) else []
    if not scenes:
        raise RuntimeError("AI video script me scenes nahi mile.")

    workdir = tempfile.mkdtemp(prefix="edusolve_video_")
    segments = []
    try:
        total = len(scenes)
        for idx, scene in enumerate(scenes, 1):
            heading = _clean_text(scene.get("heading", f"Scene {idx}")) or f"Scene {idx}"
            narration = str(scene.get("narration", "")).strip()
            if not narration:
                narration = f"Aaiye {heading} ko step by step samajhte hain."

            audio_path = os.path.join(workdir, f"audio_{idx}.mp3")
            duration = 8.0
            try:
                gTTS(text=_clean_text(narration)[:6000], lang=get_lang_code(language), slow=False).save(audio_path)
                duration_cmd = ["ffprobe","-v","error","-show_entries","format=duration","-of","default=noprint_wrappers=1:nokey=1",audio_path]
                duration = max(6.0, min(45.0, float(subprocess.check_output(duration_cmd, text=True).strip()) + 0.25))
            except Exception as e:
                print(f"TTS fallback: {e}")
                duration = max(8.0, min(30.0, 0.045 * len(_clean_text(narration)) + 2))
                _silent_audio(audio_path, duration)

            # Four visual keyframes make the scene feel animated instead of static.
            frame_paths=[]
            for k in range(4):
                frame_path=os.path.join(workdir, f"scene_{idx}_{k}.png")
                img=Image.new("RGB", (1280,720), (8,15,30))
                _draw_scene_visual(img, scene, k/3.0, k, 4, language)
                img.save(frame_path, "PNG", optimize=True)
                frame_paths.append(frame_path)

            # Hold each keyframe for an equal portion of the narration.
            frame_list=os.path.join(workdir, f"frames_{idx}.txt")
            part=duration/4.0
            with open(frame_list,"w",encoding="utf-8") as f:
                for fp in frame_paths:
                    f.write(f"file '{fp.replace(chr(39), chr(39)+chr(92)+chr(39)+chr(39))}'\n")
                    f.write(f"duration {part:.3f}\n")
                f.write(f"file '{frame_paths[-1].replace(chr(39), chr(39)+chr(92)+chr(39)+chr(39))}'\n")

            silent_segment=os.path.join(workdir, f"video_{idx}.mp4")
            cmd=["ffmpeg","-y","-f","concat","-safe","0","-i",frame_list,"-i",audio_path,"-t",f"{duration:.3f}","-r","24","-c:v","libx264","-preset","veryfast","-pix_fmt","yuv420p","-c:a","aac","-b:a","128k","-shortest",silent_segment]
            subprocess.run(cmd,check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
            segments.append(silent_segment)

        concat_file=os.path.join(workdir,"concat.txt")
        with open(concat_file,"w",encoding="utf-8") as f:
            for segment in segments:
                f.write("file '"+segment.replace("'","'\\''")+"'\n")
        output_path=os.path.join(workdir,"EduSolve_AI_Visual_Teacher_Lesson.mp4")
        subprocess.run(["ffmpeg","-y","-f","concat","-safe","0","-i",concat_file,"-c","copy",output_path],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
        if not os.path.exists(output_path) or os.path.getsize(output_path)<5000:
            raise RuntimeError("Video file create nahi hui.")
        return {"path":output_path,"title":script.get("title",topic),"scenes":total,"workdir":workdir,"provider":"EduSolve AI Visual Teacher","mode":"visual_fallback","note":"Animated teacher + progressive diagrams + narration; no paid video API required."}
    except Exception:
        shutil.rmtree(workdir, ignore_errors=True)
        raise


def _heygen_key():
    try:
        key = st.secrets.get("HEYGEN_API_KEY")
        if key:
            return str(key).strip()
    except Exception:
        pass
    return os.getenv("HEYGEN_API_KEY", "").strip()


def _heygen_headers():
    key = _heygen_key()
    if not key:
        raise RuntimeError("HEYGEN_API_KEY missing. Streamlit Secrets me HEYGEN_API_KEY add karo for real human-presenter AI videos.")
    return {"X-Api-Key": key, "Accept": "application/json", "Content-Type": "application/json"}


def _heygen_get_json(url, timeout=30):
    import requests
    r = requests.get(url, headers=_heygen_headers(), timeout=timeout)
    if r.status_code >= 400:
        try:
            detail = r.json()
        except Exception:
            detail = r.text[:800]
        raise RuntimeError(f"HeyGen API {r.status_code}: {detail}")
    return r.json()


def create_real_ai_video(topic, language="Hinglish", style="Cinematic Classroom", duration="Deep (6-10 min)", avatar_id="", voice_id="", style_id=""):
    """Generate a real presenter-led educational video through HeyGen Video Agent v3.

    The provider handles the human presenter, lip-sync, scene composition and visual
    rendering. The prompt explicitly asks for deep teaching, animated diagrams,
    examples, code/equations where useful, and frequent visual changes rather than
    a static slide deck. This is optional: the rest of EduSolve AI remains usable
    without a HeyGen key.
    """
    import requests
    topic = _clean_text(topic)
    if not topic:
        raise ValueError("Topic required")
    key = _heygen_key()
    if not key:
        raise RuntimeError("Real AI video ke liye HEYGEN_API_KEY required hai. Streamlit Secrets me key add karo.")

    duration_map = {
        "Quick (2-4 min)": "2 to 4 minutes",
        "Deep (6-10 min)": "6 to 10 minutes",
        "Masterclass (10-15 min)": "10 to 15 minutes",
    }
    dur = duration_map.get(duration, duration)
    lang_instruction = {
        "Hinglish": "natural Indian Hinglish, keeping technical terms in English",
        "Hindi": "natural Hindi, retaining standard technical terms in English where appropriate",
        "Gujarati": "natural Gujarati, retaining standard technical terms in English where appropriate",
        "Marathi": "natural Marathi, retaining standard technical terms in English where appropriate",
        "English": "clear natural English",
    }.get(language, language)

    if re.search(r"\besp32\b", topic, re.I):
        topic_lock = (
            "ESP32 TOPIC LOCK: If the topic is ESP32, cover only ESP32 and explicitly "
            "distinguish ESP32 from ESP8266, Arduino UNO and other boards. Cover what it is, "
            "architecture, GPIO, ADC/DAC, PWM, serial protocols, Wi-Fi/Bluetooth, programming, "
            "IoT example, applications, limitations and recap. Never substitute another controller. "
            "If a specification depends on the exact ESP32 variant, say so rather than guessing."
        )
    else:
        topic_lock = f"TOPIC LOCK: Stay exactly on {topic}. Do not drift to adjacent topics."

    prompt = f"""
Create a premium educational explainer video about EXACTLY this topic: {topic}

{topic_lock}

Audience: college/university students.
Language: {lang_instruction}.
Target duration: {dur}.
Visual style: {style}.

IMPORTANT: This must NOT look like a slideshow with a voiceover. Make it feel like a
modern AI educational YouTube masterclass with a realistic human presenter/teacher
who is visibly speaking and explaining. The presenter should appear throughout the
lesson, use natural gestures and camera changes, and interact with the visuals.

TEACHING DEPTH — COMPLETE LESSON CONTRACT:
The student needs the topic actually explained, not merely introduced. First classify the exact
topic and teach its canonical core concepts without drifting into unrelated topics. Build the
lesson from zero to confident understanding. Use 10-14 logical teaching beats/scenes as needed.

Cover, when applicable: why/problem and learning goal; exact definition; intuition/analogy;
prerequisites; main components; complete step-by-step working; architecture/data flow;
formula/code/algorithm representation; at least two concrete worked examples; edge cases;
comparison and common misconceptions; applications; limitations; exam/viva/practical
takeaways; and a final connected recap with 5 self-check questions.

Explain WHAT, WHY and HOW. Define technical terms before using them. Do not simply read
bullet points. For algorithms show intermediate states and complexity when applicable. For
programming show complete relevant code and explain important blocks. For mathematics show
symbols, substitution and interpretation. For electronics/IoT show components, connections
and signal/data flow. Never use '...', 'rest of code', 'insert image', or other placeholders.

VISUALIZATION RULES:
- Use dynamic diagrams, arrows, labels, callouts, highlighted objects, simple animations,
  charts/tables, code overlays, formulas and contextual visual examples whenever useful.
- Change visual composition regularly; do not keep one background slide on screen.
- Keep the human presenter visible while the visual explanation is happening.
- On-screen text must reinforce the explanation, not replace it.
- Keep the exact requested topic locked throughout the video.
- Never invent technical specifications, facts, sources or citations.
- The final scene must connect the entire lesson into one mental model and ask 5 self-check questions.

Produce a polished, coherent, fact-focused masterclass. The presenter should teach like a
patient expert, not merely read a script.

""".strip()

    payload = {
        "prompt": prompt,
        "mode": "generate",
        "orientation": "landscape",
        "incognito_mode": False,
    }
    if avatar_id.strip():
        payload["avatar_id"] = avatar_id.strip()
    if voice_id.strip():
        payload["voice_id"] = voice_id.strip()
    if style_id.strip():
        payload["style_id"] = style_id.strip()

    r = requests.post("https://api.heygen.com/v3/video-agents", headers=_heygen_headers(), json=payload, timeout=45)
    if r.status_code >= 400:
        try:
            detail = r.json()
        except Exception:
            detail = r.text[:1000]
        raise RuntimeError(f"HeyGen create failed ({r.status_code}): {detail}")
    data = r.json().get("data", r.json())
    session_id = data.get("session_id")
    video_id = data.get("video_id")
    if not session_id and not video_id:
        raise RuntimeError(f"HeyGen ne session/video id nahi diya: {r.json()}")

    deadline = time.time() + 15 * 60
    last_status = "generating"
    while time.time() < deadline:
        if video_id:
            vd = _heygen_get_json(f"https://api.heygen.com/v3/videos/{video_id}")
            vdata = vd.get("data", vd)
            status = str(vdata.get("status", "")).lower()
            last_status = status or last_status
            if status == "completed" and vdata.get("video_url"):
                video_url = vdata["video_url"]
                break
            if status == "failed":
                raise RuntimeError(f"HeyGen video failed: {vdata.get('failure_message') or vdata}")
        else:
            sd = _heygen_get_json(f"https://api.heygen.com/v3/video-agents/{session_id}")
            sdata = sd.get("data", sd)
            status = str(sdata.get("status", "")).lower()
            last_status = status or last_status
            video_id = sdata.get("video_id") or video_id
            if status == "failed":
                raise RuntimeError(f"HeyGen agent failed: {sdata.get('failure_message') or sdata}")
            if not video_id and status == "completed":
                # Some responses expose the rendered video directly.
                direct = sdata.get("video_url")
                if direct:
                    video_url = direct
                    break
        time.sleep(8)
    else:
        raise RuntimeError(f"HeyGen video timeout. Last status: {last_status}. Session: {session_id}")

    out_dir = tempfile.mkdtemp(prefix="edusolve_real_video_")
    out_path = os.path.join(out_dir, "EduSolve_AI_Real_AI_Teacher.mp4")
    rr = requests.get(video_url, timeout=120)
    if rr.status_code >= 400:
        raise RuntimeError(f"Video download failed: HTTP {rr.status_code}")
    with open(out_path, "wb") as f:
        f.write(rr.content)
    if os.path.getsize(out_path) < 10000:
        raise RuntimeError("Downloaded video file is unexpectedly small.")
    return {
        "path": out_path,
        "title": topic,
        "provider": "HeyGen Video Agent",
        "video_id": video_id or "",
        "session_id": session_id or "",
        "mode": "real_human_visual",
    }

def transcribe_audio(api_key, audio_path):
    try:
        key = api_key or get_api_key()
        if not key:
            return None

        client = Groq(api_key=key)
        with open(audio_path, "rb") as f:
            transcription = client.audio.transcriptions.create(
                file=(os.path.basename(audio_path), f.read()),
                model="whisper-large-v3-turbo",
                language="hi",
                response_format="text",
            )
        return transcription if isinstance(transcription, str) else transcription.text
    except Exception as e:
        print(f"Whisper Error: {e}")
        return None


# ============================================================
# DOCUMENT PROCESSING / RAG
# ============================================================

def process_files(files, chunk_size=1000, chunk_overlap=100):
    if not files:
        st.warning("No files uploaded")
        return

    docs = []
    processed_names = []

    for file in files:
        ext = os.path.splitext(file.name)[-1].lower()
        path = None

        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
                tmp.write(file.getvalue())
                path = tmp.name

            if ext == ".pdf":
                loader = PyPDFLoader(path)
            elif ext == ".csv":
                loader = CSVLoader(path)
            elif ext == ".txt":
                loader = TextLoader(path, encoding="utf-8")
            else:
                st.warning(f"Unsupported file skipped: {file.name}")
                continue

            loaded = loader.load()
            for d in loaded:
                d.metadata["source"] = file.name
                d.metadata["file_name"] = file.name
            clean_loaded = [d for d in loaded if d.page_content.strip()]
            docs.extend(clean_loaded)
            if clean_loaded:
                processed_names.append(file.name)

        except Exception as e:
            st.error(f"Error loading {file.name}: {e}")
        finally:
            if path and os.path.exists(path):
                try:
                    os.remove(path)
                except Exception:
                    pass

    if not docs:
        st.error("No content extracted!")
        return

    chunk_size = max(300, int(chunk_size))
    chunk_overlap = max(0, min(int(chunk_overlap), chunk_size - 1))

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    chunks = splitter.split_documents(docs)

    try:
        vectordb = FAISS.from_documents(chunks, get_embeddings())
    except Exception as e:
        st.error(f"Embedding/FAISS error: {e}")
        return

    st.session_state.vectordb = vectordb
    st.session_state.history = []
    st.session_state.weak_topics = {}
    st.session_state.kb_files = processed_names
    st.session_state.kb_chunk_count = len(chunks)

    st.success(
        f"✅ {len(chunks)} chunks ready from {len(docs)} document sections!"
    )


def hybrid_search(query, k=8):
    """Robust lightweight hybrid retrieval without an extra BM25 dependency.

    FAISS provides semantic retrieval. A small lexical bonus is then applied so
    exact terms from the student's query are less likely to disappear from the
    final context. This keeps deployment simple while improving RAG quality.
    """
    vb = st.session_state.get("vectordb")
    if vb is None:
        return []

    try:
        k = max(1, int(k))
        candidate_n = min(max(k * 3, 12), 40)
        candidates = vb.similarity_search_with_score(query, k=candidate_n)
    except Exception:
        try:
            return vb.similarity_search(query, k=k)
        except Exception as e:
            print(f"Search Error: {e}")
            return []

    terms = set(re.findall(r"[a-zA-Z0-9_]{3,}", str(query).lower()))
    scored = []
    for doc, distance in candidates:
        text = str(doc.page_content or "").lower()
        lexical = sum(text.count(term) for term in terms)
        # FAISS distance is lower-is-better for the common indexes used here.
        semantic = 1.0 / (1.0 + max(float(distance), 0.0))
        score = semantic + min(lexical, 8) * 0.035
        scored.append((score, doc))

    scored.sort(key=lambda x: x[0], reverse=True)
    return [doc for _, doc in scored[:k]]


def _context_for(query, k=8, per_doc=1600):
    docs = hybrid_search(query, k=k)
    pieces = []
    for d in docs:
        source = d.metadata.get("source", "Unknown")
        page = d.metadata.get("page")
        page_text = f" | page {page + 1}" if isinstance(page, int) else ""
        pieces.append(
            f"[SOURCE: {source}{page_text}]\n{d.page_content[:per_doc]}"
        )
    return "\n\n".join(pieces)[:MAX_CONTEXT_CHARS], docs



def generate_answer_visual(question, answer, language="Hinglish", source_context=""):
    """Create a factual visual plan for an answer.

    Returns a normalized JSON-like dict. The renderer only draws relationships/data
    explicitly returned by the model, so it does not invent graph points or wiring.
    """
    q = _clean_text(question)
    a = str(answer or "")[:12000]
    ctx = str(source_context or "")[:9000]
    fallback = {"kind": "none", "title": "", "caption": ""}
    prompt = f"""You are EduSolve AI's visual reasoning engine. Create ONE useful visual aid for this exact student question.\nQuestion: {q}\nLanguage: {language}\nAnswer: {a}\nSource context (if present): {ctx or 'None'}\n\nReturn ONLY JSON. Choose kind from: concept, flow, architecture, comparison, timeline, graph, table, none.\nRules:\n- The visual MUST explain the exact question; do not drift to a related topic.\n- For a graph, include ONLY explicit numeric points supplied by the answer/context or universally defined by a stated equation. If there is not enough numeric data, choose another kind.\n- For diagrams, provide short nodes and explicit edges.\n- For comparisons, provide aligned columns/rows.\n- For technical/IoT questions, show only components/links explicitly supported by the answer/context.\n- Do not invent pin numbers, specifications, measurements, values, or facts.\nSchema:\n{{\"kind\":\"flow\",\"title\":\"...\",\"nodes\":[\"...\"],\"edges\":[[0,1]],\"steps\":[\"...\"],\"columns\":[\"...\"],\"rows\":[[\"...\",\"...\"]],\"points\":[[0,0]],\"x_label\":\"...\",\"y_label\":\"...\",\"caption\":\"...\"}}\nUse empty arrays for unused fields."""
    try:
        data = _json_from_text(_invoke(prompt, temperature=0.15, max_tokens=2500), default=None)
        if not isinstance(data, dict):
            return fallback
        kind = str(data.get("kind") or "none").lower().strip()
        if kind not in {"concept","flow","architecture","comparison","timeline","graph","table","none"}:
            kind = "none"
        data["kind"] = kind
        data["title"] = _clean_text(data.get("title") or q[:90])
        data["caption"] = _clean_text(data.get("caption") or "")
        for key in ("nodes","steps","columns"):
            v=data.get(key,[]); data[key]=v if isinstance(v,list) else []
            data[key]=[_clean_text(x) for x in data[key]][:12]
        rows=data.get("rows",[]); data["rows"]=rows if isinstance(rows,list) else []
        data["rows"]=[([_clean_text(x) for x in r] if isinstance(r,list) else []) for r in data["rows"][:12]]
        edges=data.get("edges",[]); data["edges"]=edges if isinstance(edges,list) else []
        clean_edges=[]
        for e in data["edges"][:20]:
            if isinstance(e,list) and len(e)>=2:
                try: clean_edges.append([int(e[0]),int(e[1])])
                except Exception: pass
        data["edges"]=clean_edges
        pts=data.get("points",[]); data["points"]=pts if isinstance(pts,list) else []
        clean_pts=[]
        for pt in data["points"][:40]:
            if isinstance(pt,(list,tuple)) and len(pt)>=2:
                try: clean_pts.append([float(pt[0]),float(pt[1])])
                except Exception: pass
        data["points"]=clean_pts
        return data
    except Exception as e:
        print(f"Answer visual generation error: {e}")
        return fallback


def render_answer_visual(question, visual, out_dir=None):
    """Render an answer visual to a PNG using only the supplied visual specification."""
    if not isinstance(visual, dict) or visual.get("kind") == "none":
        return None
    from PIL import ImageDraw, ImageFont
    out_dir = out_dir or tempfile.mkdtemp(prefix="edusolve_visual_")
    os.makedirs(out_dir, exist_ok=True)
    path=os.path.join(out_dir,"answer_visual.png")
    W,H=1500,900
    img=Image.new("RGB",(W,H),(248,250,252)); d=ImageDraw.Draw(img)
    font_path=_find_font("English")
    def font(sz):
        try: return ImageFont.truetype(font_path,sz) if font_path else ImageFont.load_default()
        except Exception: return ImageFont.load_default()
    title=font(42); body=font(25); small=font(20)
    d.rounded_rectangle((35,35,W-35,H-35),30,fill=(255,255,255),outline=(203,213,225),width=3)
    d.text((70,65),"EduSolve AI • Visual Explanation",font=small,fill=(79,70,229))
    d.text((70,105),str(visual.get("title") or question)[:70],font=title,fill=(15,23,42))
    kind=visual.get("kind")
    if kind=="graph":
        try:
            import matplotlib.pyplot as plt
            pts=visual.get("points",[])
            if not pts: return None
            fig,ax=plt.subplots(figsize=(12,6),dpi=120)
            xs=[p[0] for p in pts]; ys=[p[1] for p in pts]
            ax.plot(xs,ys,marker="o",linewidth=2.5)
            ax.set_xlabel(visual.get("x_label") or "x"); ax.set_ylabel(visual.get("y_label") or "y")
            ax.set_title(visual.get("title") or question); ax.grid(True,alpha=.25)
            fig.tight_layout(); fig.savefig(path,bbox_inches="tight"); plt.close(fig); return path
        except Exception: return None
    if kind in {"comparison","table"}:
        cols=visual.get("columns",[]) or ["Point","Explanation"]
        rows=visual.get("rows",[])
        x0,y0=70,190; colw=min(1350/max(1,len(cols)),420); rh=78
        for j,c in enumerate(cols):
            x=x0+j*colw; d.rounded_rectangle((x,y0,x+colw-10,y0+rh),12,fill=(79,70,229)); d.text((x+12,y0+22),str(c)[:28],font=small,fill=(255,255,255))
        for i,row in enumerate(rows[:7]):
            for j,c in enumerate(row[:len(cols)]):
                x=x0+j*colw; y=y0+(i+1)*rh; d.rounded_rectangle((x,y,x+colw-10,y+rh-8),10,fill=(248,250,252),outline=(226,232,240)); d.text((x+12,y+17),str(c)[:34],font=small,fill=(30,41,59))
    else:
        nodes=visual.get("nodes",[]) or visual.get("steps",[])
        nodes=[str(x) for x in nodes[:8]]
        if not nodes: return None
        cols=2 if len(nodes)>4 else 1; rows=(len(nodes)+cols-1)//cols
        positions=[]
        for i,n in enumerate(nodes):
            c=i%cols; r=i//cols; x=110+c*690; y=200+r*145; positions.append((x,y))
            d.rounded_rectangle((x,y,x+570,y+100),22,fill=(239,246,255),outline=(59,130,246),width=3)
            # wrap
            words=n.split(); lines=[]; cur=""
            for w in words:
                t=(cur+" "+w).strip()
                if len(t)>38 and cur: lines.append(cur); cur=w
                else: cur=t
            if cur: lines.append(cur)
            for k,line in enumerate(lines[:3]): d.text((x+25,y+18+k*27),line,font=body,fill=(15,23,42))
        edges=visual.get("edges",[])
        for e in edges:
            if len(e)>=2 and 0<=e[0]<len(positions) and 0<=e[1]<len(positions):
                x1,y1=positions[e[0]]; x2,y2=positions[e[1]]; d.line((x1+285,y1+100,x2+285,y2),fill=(99,102,241),width=5)
    cap=visual.get("caption")
    if cap: d.text((70,H-100),str(cap)[:120],font=small,fill=(71,85,105))
    img.save(path,quality=95)
    return path

def generate_general_ai_answer(question, language="Hinglish", subject="", depth="Deep + Exam Ready"):
    """Answer arbitrary student questions without requiring an uploaded source.

    This is intentionally separate from RAG so the user can use EduSolve AI as
    a general teaching assistant. The prompt asks for complete explanations,
    examples, code where relevant, mistakes, exam points and a recap.
    """
    question = _clean_text(question)
    if not question:
        return "Please question type karo."

    subject_hint = _clean_text(subject) or "Not specified"
    if depth == "Standard":
        length_instruction = "Give a clear medium-length answer with the essential reasoning and one example."
    elif depth == "Detailed":
        length_instruction = "Give a detailed answer with step-by-step explanation, examples and common mistakes."
    else:
        length_instruction = (
            "Give a deep, long, complete student-ready answer. Do not skip important steps. "
            "If the question asks for code, provide the COMPLETE runnable code in one or more fenced code blocks; "
            "never use '...', 'rest of code', placeholders, or truncated sections. Explain the code line-by-line or section-by-section."
        )

    prompt = f"""You are EduSolve AI, an expert educational tutor for college and diploma students.

Student language: {language}
Optional subject: {subject_hint}
Student question: {question}
Answer depth: {depth}

{length_instruction}

Rules:
- You do NOT need any uploaded PDF or source for this answer.
- Answer the actual question directly and completely.
- Be technically accurate and honest about uncertainty. Never invent citations, experiments, specifications or facts.
- Use simple language first, then technical depth.
- For numerical problems, show formulas, substitutions, calculations and final answer.
- For programming/electronics/IoT questions, include assumptions, wiring/pin tables when useful, complete code when requested or clearly useful, setup steps, testing and troubleshooting.
- For conceptual questions, include definition, intuition, working, example, advantages/limitations, applications and exam-ready points when relevant.
- For comparison questions, use a table.
- For 'how to' questions, give ordered steps.
- Do not ask the student to upload a PDF just to answer a general question.
- End with 'Quick Revision' containing 5-8 concise takeaways and 'Exam Tip' containing practical advice.

Return a polished answer with clear Markdown headings.
"""
    try:
        answer = _invoke_long(prompt, temperature=0.25, max_tokens=7500)
        if answer:
            return answer
    except Exception as e:
        return f"## AI Mode Error\n\n{e}"
    return "AI answer generate nahi ho saka. Please dobara try karo."


def ask_question(query, k=8, mode="normal", language="Hinglish"):
    vb = st.session_state.get("vectordb")
    if vb is None:
        return "Pehle File Upload karo.", []

    docs = hybrid_search(query, k=k)
    if not docs:
        return "PDF me ye topic nahi mila.", []

    ctx_parts = []
    for d in docs:
        source = d.metadata.get("source", "Unknown")
        page = d.metadata.get("page")
        page_text = f" page {page + 1}" if isinstance(page, int) else ""
        ctx_parts.append(
            f"[SOURCE: {source}{page_text}]\n{d.page_content[:1800]}"
        )
    ctx = "\n\n".join(ctx_parts)[:MAX_CONTEXT_CHARS]

    prompt = f"""
You are a helpful AI teaching assistant.
Language: {language}
Student question: {query}

Answer using the supplied study context. Keep the explanation student-friendly.
Use headings/bullets when useful. If a claim is not supported by the context,
say 'Context me nahi hai' rather than inventing it.
At the end, add a short 'Quick Recall' section with 2-4 key points.

STUDY CONTEXT:
{ctx}
"""

    try:
        ans = _invoke(prompt, temperature=0.2)
    except Exception as e:
        return f"AI error: {e}", [d.metadata for d in docs]

    if "history" not in st.session_state:
        st.session_state.history = []

    st.session_state.history.append(
        {
            "query": query,
            "answer": ans,
            "sources": [d.metadata for d in docs],
        }
    )
    return ans, [d.metadata for d in docs]


# ============================================================
# QUIZ / PRACTICE
# ============================================================

def generate_quiz(num_q=5, language="Hinglish", topic=None):
    topic_text = topic or "the uploaded study material"
    search_query = topic or "important concepts definitions"
    docs = hybrid_search(search_query, k=8)

    if not docs:
        return [
            {
                "question": f"No study material found for {topic_text}. What should you do first?",
                "options": [
                    "a) Upload study material",
                    "b) Delete the app",
                    "c) Skip studying",
                    "d) None",
                ],
                "answer": "a) Upload study material",
                "explanation": "Upload the relevant PDF/TXT/CSV first.",
                "topic": "setup",
            }
        ]

    ctx = "\n".join(
        f"[SOURCE {d.metadata.get('source','Unknown')}] {d.page_content[:1200]}"
        for d in docs
    )[:8000]

    prompt = f"""
Create {int(num_q)} high-quality MCQs for a student.
Language: {language}
Topic focus: {topic_text}

Use ONLY this context:
{ctx}

Return ONLY a JSON array. Every item must contain:
question, options (exactly 4 strings), answer (one exact option string),
explanation, topic.
Make distractors plausible and do not repeat the same question.
"""

    try:
        data = _json_from_text(_invoke(prompt, temperature=0.4), default=None)
        if isinstance(data, list) and data:
            cleaned = []
            for item in data[: int(num_q)]:
                if isinstance(item, dict):
                    options = item.get("options", [])
                    if not isinstance(options, list):
                        options = []
                    options = [str(x).strip() for x in options[:4] if str(x).strip()]
                    answer = str(item.get("answer", "")).strip()
                    question = str(item.get("question", "")).strip()
                    if question and len(options) == 4 and answer:
                        cleaned.append(
                            {
                                "question": question,
                                "options": options,
                                "answer": answer,
                                "explanation": str(item.get("explanation", "Review the study context.")).strip(),
                                "topic": str(item.get("topic", topic_text)).strip() or topic_text,
                            }
                        )
            if cleaned:
                return cleaned
    except Exception as e:
        print(f"Quiz generation error: {e}")

    return [
        {
            "question": f"What is an important concept from {topic_text}?",
            "options": [
                "a) Refer to the uploaded notes",
                "b) Unrelated option",
                "c) Unrelated option",
                "d) Unrelated option",
            ],
            "answer": "a) Refer to the uploaded notes",
            "explanation": "Review the uploaded material for the exact concept.",
            "topic": topic_text,
        }
        for _ in range(int(num_q))
    ]


def check_quiz_answer(question, user_ans, correct_ans, topic="general"):
    normalize = lambda x: re.sub(r"\s+", " ", str(x).strip().lower())
    a = normalize(user_ans)
    b = normalize(correct_ans)

    # Exact text first; then compare option letters such as a), b), etc.
    ok = a == b
    if not ok:
        ma = re.match(r"^([a-d])(?:[\).:\-\s]|$)", a)
        mb = re.match(r"^([a-d])(?:[\).:\-\s]|$)", b)
        ok = bool(ma and mb and ma.group(1) == mb.group(1))

    if not ok:
        wt = st.session_state.get("weak_topics", {})
        wt[topic] = wt.get(topic, 0) + 1
        st.session_state.weak_topics = wt

    return ok, ("✅ Sahi Jawaab!" if ok else f"❌ Galat. Sahi: {correct_ans}")


def build_quiz_final_report(quiz_data, answers, language="Hinglish"):
    """Create a deterministic final quiz report with every wrong answer and its correct answer."""
    report=[]
    correct=0
    for i,q in enumerate(quiz_data or []):
        user_ans = str((answers or {}).get(i, "")).strip()
        correct_ans = str(q.get("answer", "")).strip()
        ok,_ = check_quiz_answer(q.get("question", ""), user_ans, correct_ans, q.get("topic", "general"))
        if ok: correct += 1
        report.append({
            "number": i+1,
            "question": q.get("question", ""),
            "user_answer": user_ans or "Not answered",
            "correct_answer": correct_ans,
            "is_correct": bool(ok),
            "explanation": q.get("explanation", ""),
            "topic": q.get("topic", "General"),
        })
    total=len(report)
    return {"total":total,"correct":correct,"wrong":total-correct,"score":round(correct/total*100) if total else 0,"items":report}


def build_viva_final_report(viva_data, evaluations, language="Hinglish"):
    """Create a final viva report including student answer, score, correct/reference answer and missing points."""
    items=[]
    total=len(viva_data or [])
    earned=0.0
    for i,q in enumerate(viva_data or []):
        ev=(evaluations or {}).get(i,{})
        try: score=float(ev.get("score",0))
        except Exception: score=0.0
        earned += max(0,min(10,score))
        items.append({
            "number":i+1,
            "question":q.get("q", ""),
            "user_answer":ev.get("user_answer", "Not answered"),
            "correct_answer":q.get("a", ""),
            "score":round(score,1),
            "is_correct":bool(ev.get("is_correct",False)),
            "feedback":ev.get("feedback", ""),
            "missing_points":ev.get("missing_points",[]) or [],
        })
    percent=round((earned/(total*10))*100) if total else 0
    return {"total":total,"earned":round(earned,1),"score":percent,"items":items}


def generate_adaptive_quiz(num_q=5, language="Hinglish", weak_topics=None):
    """Generate practice biased toward the student's weakest tracked topics."""
    weak_topics = weak_topics or get_weak_topics()
    if weak_topics:
        focus = ", ".join(
            t for t, _ in sorted(weak_topics.items(), key=lambda x: x[1], reverse=True)[:5]
        )
    else:
        focus = "core concepts from the uploaded material"
    return generate_quiz(num_q=num_q, language=language, topic=focus)


def get_weak_topics():
    return st.session_state.get("weak_topics", {})


# ============================================================
# STUDY CONTENT GENERATORS
# ============================================================

def generate_summary(language="Hinglish", topic=None):
    query = topic or "summary overview"
    ctx, _ = _context_for(query, k=10, per_doc=1500)
    if not ctx:
        return "Pehle study material upload karo."

    prompt = f"""
Create a detailed but easy-to-revise study summary.
Language: {language}
Topic: {topic or 'all uploaded material'}
Use ONLY the context below.
Include: core concepts, definitions, important relationships, examples,
and a final 10-point quick revision list.

CONTEXT:
{ctx}
"""
    try:
        return _invoke(prompt, temperature=0.2)
    except Exception as e:
        return f"AI error: {e}"


def predict_important_questions(language="Hinglish", topic=None):
    query = topic or "exam important questions"
    ctx, _ = _context_for(query, k=10, per_doc=1500)
    if not ctx:
        return "Pehle study material upload karo."

    prompt = f"""
Based ONLY on the supplied study material, create a revision-oriented list of
likely important questions. Do NOT claim that these questions are guaranteed
to appear in an exam.
Language: {language}
Topic: {topic or 'uploaded material'}
For each question, include a short reason based on the material and a concise
answer outline.

CONTEXT:
{ctx}
"""
    try:
        return _invoke(prompt, temperature=0.3)
    except Exception as e:
        return f"AI error: {e}"


def load_conversation_history():
    return st.session_state.get("history", [])


def story_mode_learning(topic, language="Hinglish"):
    ctx, _ = _context_for(topic, k=8, per_doc=1400)
    if not ctx:
        return "Pehle relevant study material upload karo."

    prompt = f"""
Teach {topic} as a memorable short story for a student.
Language: {language}
Use ONLY the study context. Preserve technical correctness.
After the story, add 'Real Meaning' and '3 Things To Remember'.

CONTEXT:
{ctx}
"""
    try:
        return _invoke(prompt, temperature=0.8)
    except Exception as e:
        return f"AI error: {e}"


def build_project_guide(idea, budget="low", language="Hinglish"):
    """Generate a complete practical project guide, with special care for IoT code.

    The old single short completion could truncate Arduino code. This version
    allocates a long completion and explicitly forbids ellipses/placeholders.
    """
    idea = _clean_text(idea) or "IoT project"
    ctx, _ = _context_for(idea, k=8, per_doc=1500)

    prompt = f"""You are an expert IoT project mentor for Indian diploma/college students.
Project idea: {idea}
Budget: {budget}
Language: {language}

Relevant uploaded context (optional):
{ctx or 'No relevant uploaded context. Use general engineering knowledge and label assumptions clearly.'}

Create a COMPLETE build guide. It must be practical enough for a student to build and test the project.

Required sections:
1. Project overview and working principle
2. Components table — exact component, quantity, purpose, and approximate INR range
3. Complete wiring/pin table — board pin, component pin, power, ground, notes
4. POWER PLAN — voltage/current considerations and safe wiring
5. COMPLETE FIRMWARE CODE — if ESP32/Arduino/microcontroller is involved, provide the ENTIRE code in one or more fenced code blocks.
   - Never write '...', 'rest of code', 'same as above', 'add your code here', or omit functions.
   - Include all imports/includes, definitions, setup, loop, helper functions and configuration.
   - Code must be internally consistent with the wiring table.
6. Software/app setup — Arduino IDE, libraries, Blynk or other app if applicable
7. Step-by-step assembly
8. Testing procedure with expected results
9. Troubleshooting table
10. Total budget estimate
11. Viva questions and answers
12. Future enhancements

For uncertain hardware details, state the assumption and give a safe example instead of pretending it is exact.
For Blynk projects, include template/device setup, virtual pins and dashboard controls when relevant.
For ESP32 projects, prefer GPIO numbers and explain board-label equivalents only when known.
Do not shorten the answer merely to save space.
"""
    try:
        answer = _invoke_long(prompt, temperature=0.35, max_tokens=8000)
        if answer:
            return answer
    except Exception as e:
        return f"## Project Guide Error\n\n{e}"
    return "Complete project guide generate nahi ho saka. Please dobara try karo."


def generate_podcast_script(topic, language="Hinglish"):
    ctx, _ = _context_for(topic, k=6, per_doc=1300)
    if not ctx:
        return "Pehle relevant study material upload karo."

    prompt = f"""
Write an engaging 2-host educational podcast about {topic}.
Language: {language}
Hosts: Host 1 and Host 2.
Use ONLY the context below.
Include intro, explanation, example, common mistake, recap, and a short quiz.

CONTEXT:
{ctx}
"""
    try:
        return _invoke(prompt, temperature=0.8)
    except Exception as e:
        return f"AI error: {e}"


def generate_viva_questions(topic, num_q=5, language="Hinglish"):
    ctx, _ = _context_for(topic, k=8, per_doc=1300)
    if not ctx:
        return [{"q": f"Explain {topic}.", "a": "Upload study material for a source-grounded answer."}]

    prompt = f"""
Generate {int(num_q)} viva questions for {topic}.
Language: {language}
Use ONLY the context.
Return ONLY JSON array with objects: {{"q":"...","a":"..."}}
Mix definition, why/how, application, comparison, and troubleshooting questions.

CONTEXT:
{ctx}
"""
    try:
        data = _json_from_text(_invoke(prompt, temperature=0.4), default=None)
        if isinstance(data, list) and data:
            return [
                {"q": str(x.get("q", "")), "a": str(x.get("a", ""))}
                for x in data[: int(num_q)]
                if isinstance(x, dict)
            ]
    except Exception as e:
        print(f"Viva generation error: {e}")

    return [
        {"q": f"Explain {topic}.", "a": f"Review the uploaded material for {topic}."}
        for _ in range(int(num_q))
    ]


def verify_viva_answer(question, correct_ans, user_ans, language="Hinglish"):
    prompt = f"""
Evaluate a student's viva answer.
Language: {language}
Question: {question}
Reference answer: {correct_ans}
Student answer: {user_ans}

Return ONLY JSON:
{{"verdict":"True/False","score":0,"feedback":"...","missing_points":["..."]}}
Score from 0 to 10 based on correctness and coverage.
Do not reward confident wording if the technical content is wrong.
"""
    try:
        data = _json_from_text(_invoke(prompt, temperature=0.2), default=None)
        if isinstance(data, dict):
            data.setdefault("verdict", "False")
            data.setdefault("score", 0)
            data.setdefault("feedback", "")
            data.setdefault("missing_points", [])
            return data
    except Exception as e:
        print(f"Viva verification error: {e}")

    a = _clean_text(user_ans).lower()
    b = _clean_text(correct_ans).lower()
    ok = bool(a and (a in b or b in a))
    return {
        "verdict": str(ok),
        "score": 10 if ok else 0,
        "feedback": f"Correct answer: {correct_ans}",
        "missing_points": [],
    }


# ============================================================
# ADVANCED RETENTION HELPERS
# These are optional additions for the V24 UI and do not replace
# any original function.
# ============================================================

def generate_teach_back_feedback(topic, student_explanation, language="Hinglish"):
    ctx, _ = _context_for(topic, k=8, per_doc=1300)
    prompt = f"""
You are a strict but encouraging teacher.
Topic: {topic}
Language: {language}
Student's teach-back explanation:
{student_explanation}

Reference study context:
{ctx or 'No uploaded context available.'}

Evaluate the student's explanation against the context.
Return ONLY JSON:
{{
  "score": 0,
  "verdict": "Strong/Needs Work",
  "what_was_correct": ["..."],
  "missing_or_wrong": ["..."],
  "one_better_explanation": "...",
  "next_question": "..."
}}
Score 0-10. Do not invent facts outside the context.
"""
    try:
        data = _json_from_text(_invoke(prompt, temperature=0.2), default=None)
        if isinstance(data, dict):
            return data
    except Exception as e:
        print(f"Teach-back error: {e}")
    return {
        "score": 0,
        "verdict": "Needs Work",
        "what_was_correct": [],
        "missing_or_wrong": ["AI evaluation failed."],
        "one_better_explanation": "Try again after checking your notes.",
        "next_question": topic,
    }


def generate_revision(topic, language="Hinglish", focus="weakness"):
    ctx, _ = _context_for(topic, k=8, per_doc=1400)
    prompt = f"""
Create a compact smart-revision session for topic: {topic}
Language: {language}
Focus: {focus}
Use ONLY this context:
{ctx or 'No context available.'}

Return sections:
1. 5-minute recap
2. 5 flashcards (Q/A)
3. 3 common traps
4. 3 self-test questions without answers first
5. answer key
"""
    try:
        return _invoke(prompt, temperature=0.35)
    except Exception as e:
        return f"AI error: {e}"


def generate_exam_attack(topic, language="Hinglish"):
    """Build a reliable, source-grounded Exam Attack plan.

    The old version always called the LLM even when no source was indexed and
    returned a raw API error.  This version first verifies the knowledge base,
    then uses grounded context, and finally provides a useful deterministic
    fallback if the model call fails.
    """
    topic = (str(topic or "uploaded material").strip() or "uploaded material")
    if not _has_source_material():
        return (
            "## 🚨 Exam Attack\n\n"
            "Pehle Knowledge Base me PDF/TXT/CSV upload karo aur process hone do. "
            "Uske baad Exam Attack dobara run karo."
        )

    ctx, docs = _context_for(topic, k=10, per_doc=1400)
    if not ctx:
        return (
            "## 🚨 Exam Attack\n\n"
            f"**{topic}** ke liye uploaded material me relevant content nahi mila. "
            "Topic ka naam exactly notes ke according try karo."
        )

    prompt = f"""
You are an exam-preparation teacher. Build a practical exam-focused study
attack plan for: {topic}.
Language: {language}.
Use ONLY the supplied study material. Never claim that a question is
guaranteed to appear in an actual exam.

Return these exact sections:
1. MUST KNOW
2. DEFINITIONS TO MEMORIZE
3. COMPARISONS / DIFFERENCES
4. DIAGRAMS / PROCESSES TO PRACTICE
5. 5 SELF-TEST QUESTIONS
6. COMMON MISTAKES
7. 30-MINUTE ATTACK PLAN
8. LAST-MINUTE CHECKLIST

Keep it concise, exam-oriented and source-grounded.

SOURCE CONTEXT:
{ctx}
"""
    try:
        answer = _invoke(prompt, temperature=0.25)
        if answer and str(answer).strip():
            return str(answer).strip()
    except Exception as e:
        print(f"Exam Attack LLM error: {e}")

    # Deterministic fallback: the feature still works even if the model/API
    # temporarily fails. It never invents subject facts.
    snippets = []
    for d in docs[:8]:
        text = _clean_text(getattr(d, "page_content", ""))
        if text:
            snippets.append(text[:420])
    source_notes = "\n".join(f"- {x}" for x in snippets[:8])
    return f"""## 🚨 Exam Attack — {topic}

### 1. MUST KNOW
Use the following source-grounded notes as your revision base:
{source_notes or '- Review the uploaded material carefully.'}

### 2. DEFINITIONS TO MEMORIZE
- Mark every definition explicitly given in your uploaded notes.

### 3. COMPARISONS / DIFFERENCES
- Identify concepts in the notes that are presented together or contrasted.

### 4. DIAGRAMS / PROCESSES TO PRACTICE
- Practice every diagram, flow, architecture or step-by-step process present in the source.

### 5. 5 SELF-TEST QUESTIONS
1. Explain the main concept of {topic}.
2. Write the important definitions from {topic}.
3. Compare the related concepts in your notes.
4. Explain one process/diagram from the material.
5. What common mistake can occur while answering {topic}?

### 6. COMMON MISTAKES
- Do not add facts that are not in the uploaded material.
- Read definitions carefully and include required keywords.
- For comparisons, write both sides clearly.

### 7. 30-MINUTE ATTACK PLAN
- 10 min: Read definitions and core concepts.
- 10 min: Practice comparisons, diagrams and processes.
- 5 min: Answer the 5 self-test questions.
- 5 min: Recheck weak areas from your notes.

### 8. LAST-MINUTE CHECKLIST
- [ ] Definitions revised
- [ ] Important concepts revised
- [ ] Comparisons revised
- [ ] Diagrams/processes practiced
- [ ] Self-test completed

> AI plan generation was temporarily unavailable, so this fallback is built directly from your uploaded source context.
"""


def generate_confusion_battle(topic_a, topic_b, language="Hinglish"):
    ctx_a, _ = _context_for(topic_a, k=5, per_doc=1100)
    ctx_b, _ = _context_for(topic_b, k=5, per_doc=1100)
    prompt = f"""
Compare two study concepts for a student.
Concept A: {topic_a}
Concept B: {topic_b}
Language: {language}
Use only supplied context. Clearly mark if either concept lacks context.
Include: definition, key difference table, when used, example, memory trick,
and 5 challenge questions.

A CONTEXT:
{ctx_a or 'Not found'}

B CONTEXT:
{ctx_b or 'Not found'}
"""
    try:
        return _invoke(prompt, temperature=0.35)
    except Exception as e:
        return f"AI error: {e}"


# ============================================================
# PPT
# ============================================================


def create_ppt_file(topic, num_slides=10, language="Hinglish"):
    """Create a visual, teaching-first PPT with diagrams, flows, examples and recap.

    The generator asks the LLM for a structured storyboard, then renders the visuals
    locally with python-pptx so the resulting .pptx contains real editable shapes,
    connectors and text rather than a wall of bullets.
    """
    from pptx import Presentation
    from pptx.util import Inches as PptInches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    topic = _clean_text(topic) or "General Topic"
    n = max(5, min(25, int(num_slides)))
    ctx, _ = _context_for(topic, k=14, per_doc=1400)
    source_mode = bool(ctx and "No indexed source context" not in ctx)

    prompt = f"""
You are an expert instructional designer creating a beautiful university-level PPT.
Topic: {topic}
Language: {language}
Number of content slides: {n}
Source material: {ctx[:MAX_CONTEXT_CHARS] if source_mode else 'No uploaded source. Use accurate general academic knowledge.'}

Return ONLY valid JSON array. Each slide object MUST have:
{{
  "title": "short title",
  "type": "concept|flow|comparison|example|architecture|code|process|formula|recap",
  "explanation": "2-4 sentence teacher explanation",
  "points": ["short point", "short point", "short point"],
  "diagram": {{"center":"...","nodes":["..."],"steps":["..."]}},
  "example": "optional concrete example",
  "code": "optional complete short code example, empty string if not needed"
}}

Design rules:
- Do NOT make every slide a bullet list.
- Every slide must teach one clear idea.
- Use flow/process/architecture diagrams whenever relationships or steps exist.
- Use comparison slides for contrasting concepts.
- Use example slides with an input -> process -> output explanation.
- Use code slides only when code materially helps; code must be complete for the shown example.
- Include definitions, intuition, working, a practical example, common mistakes, applications,
  and a final revision/recap slide.
- Keep text concise enough to fit; the visual diagram should carry part of the explanation.
- Never invent source-specific facts when source material is supplied.
""".strip()

    fallback = []
    for i in range(n):
        fallback.append({
            "title": f"{topic} — Concept {i+1}",
            "type": "concept",
            "explanation": f"Understand the key idea of {topic} and how it connects to the rest of the topic.",
            "points": ["Core definition", "Important idea", "Practical relevance"],
            "diagram": {"center": topic, "nodes": ["Definition", "Working", "Example"]},
            "example": "Use a simple classroom example to connect theory with practice.",
            "code": "",
        })
    try:
        raw = _invoke_long(prompt, temperature=0.35, max_tokens=9000)
        slides = _json_from_text(raw, default=None)
        if not isinstance(slides, list) or not slides:
            slides = fallback
    except Exception as e:
        print(f"PPT generation error: {e}")
        slides = fallback

    # Normalize malformed model output.
    norm = []
    for i, sl in enumerate(slides[:n]):
        if not isinstance(sl, dict):
            continue
        points = sl.get("points", [])
        if isinstance(points, str): points = [points]
        if not isinstance(points, list): points = []
        diagram = sl.get("diagram", {})
        if not isinstance(diagram, dict): diagram = {}
        nodes = diagram.get("nodes", [])
        steps = diagram.get("steps", [])
        if isinstance(nodes, str): nodes = [nodes]
        if isinstance(steps, str): steps = [steps]
        norm.append({
            "title": _clean_text(sl.get("title") or f"{topic} — Slide {i+1}"),
            "type": str(sl.get("type") or "concept").lower(),
            "explanation": _clean_text(sl.get("explanation") or ""),
            "points": [_clean_text(x) for x in points if _clean_text(x)][:5],
            "diagram": {
                "center": _clean_text(diagram.get("center") or topic),
                "nodes": [_clean_text(x) for x in nodes if _clean_text(x)][:6],
                "steps": [_clean_text(x) for x in steps if _clean_text(x)][:6],
            },
            "example": _clean_text(sl.get("example") or ""),
            "code": str(sl.get("code") or ""),
        })
    slides = norm or fallback

    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    NAVY = RGBColor(15, 23, 42)
    BLUE = RGBColor(37, 99, 235)
    CYAN = RGBColor(6, 182, 212)
    PURPLE = RGBColor(124, 58, 237)
    GREEN = RGBColor(16, 185, 129)
    ORANGE = RGBColor(245, 158, 11)
    RED = RGBColor(239, 68, 68)
    INK = RGBColor(30, 41, 59)
    MUTED = RGBColor(71, 85, 105)
    BG = RGBColor(248, 250, 252)
    WHITE = RGBColor(255, 255, 255)
    LIGHT = RGBColor(226, 232, 240)

    def add_text(slide, text, x, y, w, h, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
        tf = box.text_frame
        tf.clear(); tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        p = tf.paragraphs[0]; p.text = str(text); p.alignment = align
        p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
        return box

    def add_card(slide, x, y, w, h, fill=WHITE, line=LIGHT, radius=True):
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                     PptInches(x), PptInches(y), PptInches(w), PptInches(h))
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
        shp.line.color.rgb = line
        return shp

    def add_arrow(slide, x1, y1, x2, y2, color=BLUE):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, PptInches(x1), PptInches(y1), PptInches(x2), PptInches(y2))
        line.line.color.rgb = color; line.line.width = Pt(2.5)
        try: line.line.end_arrowhead = True
        except Exception: pass
        return line

    def add_footer(slide, idx):
        add_text(slide, f"EduSolve AI  •  {topic}  •  {idx}", 0.55, 7.12, 12.1, 0.22, 9, False, MUTED)

    # Title / cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    add_text(slide, "EDUSOLVE AI", 0.75, 0.75, 4, 0.35, 14, True, CYAN)
    add_text(slide, topic, 0.75, 1.55, 11.7, 1.6, 40, True, WHITE)
    add_text(slide, "Visual Learning Presentation", 0.78, 3.25, 8.5, 0.55, 22, False, RGBColor(203,213,225))
    add_card(slide, 0.8, 4.25, 3.15, 1.25, fill=RGBColor(30,41,59), line=RGBColor(51,65,85))
    add_text(slide, "CONCEPT", 1.05, 4.55, 1.2, 0.3, 12, True, CYAN)
    add_text(slide, "Understand the why", 1.05, 4.88, 2.4, 0.35, 16, True, WHITE)
    add_card(slide, 4.25, 4.25, 3.15, 1.25, fill=RGBColor(30,41,59), line=RGBColor(51,65,85))
    add_text(slide, "VISUAL", 4.5, 4.55, 1.2, 0.3, 12, True, PURPLE)
    add_text(slide, "See how it works", 4.5, 4.88, 2.4, 0.35, 16, True, WHITE)
    add_card(slide, 7.7, 4.25, 3.15, 1.25, fill=RGBColor(30,41,59), line=RGBColor(51,65,85))
    add_text(slide, "PRACTICE", 7.95, 4.55, 1.2, 0.3, 12, True, GREEN)
    add_text(slide, "Apply + revise", 7.95, 4.88, 2.4, 0.35, 16, True, WHITE)
    add_text(slide, f"Language: {language}", 0.8, 6.45, 4, 0.3, 12, False, RGBColor(148,163,184))

    for i, sl in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
        # top accent
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0), PptInches(0), PptInches(13.333), PptInches(0.16))
        bar.fill.solid(); bar.fill.fore_color.rgb = [BLUE, PURPLE, CYAN, GREEN, ORANGE][(i-1) % 5]; bar.line.fill.background()
        add_text(slide, f"{i:02d}", 0.55, 0.45, 0.55, 0.45, 15, True, BLUE)
        add_text(slide, sl["title"], 1.15, 0.37, 11.45, 0.7, 25, True, INK)

        # Left: explanation + points
        add_card(slide, 0.55, 1.35, 4.05, 5.35, fill=WHITE, line=LIGHT)
        add_text(slide, "TEACHER EXPLANATION", 0.82, 1.65, 3.45, 0.3, 11, True, PURPLE)
        explanation = sl["explanation"] or "Understand the concept step by step using the visual on the right."
        add_text(slide, explanation, 0.82, 2.05, 3.45, 1.55, 16, False, INK)
        y = 3.78
        for ptxt in sl["points"][:4]:
            add_text(slide, "•", 0.85, y, 0.22, 0.3, 15, True, BLUE)
            add_text(slide, ptxt, 1.08, y, 3.05, 0.55, 13, False, MUTED)
            y += 0.68
        if sl["example"]:
            add_card(slide, 0.82, 6.0, 3.5, 0.55, fill=RGBColor(239,246,255), line=RGBColor(191,219,254))
            add_text(slide, "Example: " + sl["example"][:210], 1.0, 6.12, 3.15, 0.35, 10, False, INK)

        # Right visual canvas
        add_card(slide, 4.85, 1.35, 7.95, 5.35, fill=WHITE, line=LIGHT)
        typ = sl["type"]
        d = sl["diagram"]
        nodes = d.get("nodes", [])
        steps = d.get("steps", [])

        if typ in ("flow", "process", "architecture") or steps:
            items = steps or nodes or sl["points"][:5]
            items = items[:6]
            count = len(items)
            if count == 0: items = [topic]; count = 1
            box_w = 1.65
            gap = 0.25
            total_w = count * box_w + (count-1) * gap
            start_x = 5.15 + max(0, (7.35-total_w)/2)
            y1 = 3.05
            for j, item in enumerate(items):
                x = start_x + j*(box_w+gap)
                fill = [RGBColor(239,246,255), RGBColor(245,243,255), RGBColor(236,253,245), RGBColor(255,247,237), RGBColor(254,242,242), RGBColor(240,249,255)][j%6]
                linec = [RGBColor(147,197,253), RGBColor(196,181,253), RGBColor(110,231,183), RGBColor(253,186,116), RGBColor(252,165,165), RGBColor(125,211,252)][j%6]
                add_card(slide, x, y1, box_w, 1.35, fill=fill, line=linec)
                add_text(slide, item[:75], x+0.12, y1+0.28, box_w-0.24, 0.75, 13, True, INK, PP_ALIGN.CENTER)
                if j < count-1: add_arrow(slide, x+box_w, y1+0.68, x+box_w+gap, y1+0.68, color=BLUE)
            add_text(slide, "PROCESS / RELATIONSHIP", 5.2, 1.75, 3.2, 0.3, 11, True, BLUE)
        elif typ == "comparison":
            left = nodes[0] if nodes else "Concept A"
            right = nodes[1] if len(nodes)>1 else "Concept B"
            add_card(slide, 5.25, 2.05, 3.15, 3.75, fill=RGBColor(239,246,255), line=RGBColor(147,197,253))
            add_card(slide, 9.0, 2.05, 3.15, 3.75, fill=RGBColor(245,243,255), line=RGBColor(196,181,253))
            add_text(slide, left[:45], 5.55, 2.35, 2.55, 0.55, 18, True, BLUE, PP_ALIGN.CENTER)
            add_text(slide, right[:45], 9.3, 2.35, 2.55, 0.55, 18, True, PURPLE, PP_ALIGN.CENTER)
            lp = sl["points"][:3]; rp = sl["points"][3:6] or sl["points"][:3]
            for j, t in enumerate(lp): add_text(slide, "• "+t, 5.55, 3.2+j*0.65, 2.55, 0.5, 12, False, INK)
            for j, t in enumerate(rp): add_text(slide, "• "+t, 9.3, 3.2+j*0.65, 2.55, 0.5, 12, False, INK)
            add_text(slide, "VS", 8.35, 3.65, 0.55, 0.45, 16, True, ORANGE, PP_ALIGN.CENTER)
        elif typ == "code" and sl["code"]:
            add_text(slide, "COMPLETE WORKING EXAMPLE", 5.25, 1.75, 3.6, 0.3, 11, True, GREEN)
            add_card(slide, 5.2, 2.15, 7.25, 3.85, fill=RGBColor(15,23,42), line=RGBColor(51,65,85))
            code = sl["code"][:3200]
            add_text(slide, code, 5.48, 2.42, 6.7, 3.25, 11, False, RGBColor(226,232,240))
            add_text(slide, "What the code is doing", 5.25, 6.18, 2.5, 0.3, 11, True, PURPLE)
            add_text(slide, sl["explanation"][:260], 7.35, 6.05, 4.8, 0.55, 10, False, MUTED)
        elif typ == "formula":
            add_text(slide, d.get("center") or topic, 5.35, 2.0, 7.0, 0.65, 25, True, BLUE, PP_ALIGN.CENTER)
            add_text(slide, "↓", 8.2, 2.75, 0.5, 0.5, 24, True, PURPLE, PP_ALIGN.CENTER)
            for j, item in enumerate((nodes or sl["points"])[:4]):
                x = 5.35 + (j%2)*3.65; y = 3.45 + (j//2)*1.25
                add_card(slide, x, y, 3.15, 0.95, fill=RGBColor(239,246,255), line=RGBColor(147,197,253))
                add_text(slide, item, x+0.15, y+0.22, 2.85, 0.45, 13, True, INK, PP_ALIGN.CENTER)
        else:
            center = d.get("center") or topic
            add_card(slide, 7.2, 2.85, 3.25, 1.1, fill=RGBColor(239,246,255), line=RGBColor(96,165,250))
            add_text(slide, center[:55], 7.4, 3.14, 2.85, 0.45, 17, True, BLUE, PP_ALIGN.CENTER)
            around = nodes or sl["points"][:6]
            positions = [(5.2,2.0),(9.9,2.0),(5.15,4.75),(9.95,4.75),(6.25,5.65),(8.9,5.65)]
            for j,item in enumerate(around[:6]):
                x,y = positions[j]
                add_card(slide,x,y,2.15,0.72,fill=WHITE,line=LIGHT)
                add_text(slide,item[:38],x+0.08,y+0.15,1.99,0.38,10,True,INK,PP_ALIGN.CENTER)
                add_arrow(slide,8.82,3.4,x+1.07,y+0.02 if y<3 else y+0.02,color=RGBColor(148,163,184))

        add_footer(slide, i)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    tmp.close()
    prs.save(tmp.name)
    return tmp.name

# ============================================================
# IMAGE EXPORTS
# ============================================================

def images_to_pdf(image_files):
    pil_images = []

    for img_file in image_files:
        try:
            img_file.seek(0)
            pil_images.append(Image.open(img_file).convert("RGB"))
        except Exception:
            try:
                img_file.seek(0)
                pil_images.append(Image.open(io.BytesIO(img_file.read())).convert("RGB"))
            except Exception:
                continue

    pdf_buffer = io.BytesIO()
    if not pil_images:
        pdf_buffer.seek(0)
        return pdf_buffer

    if len(pil_images) == 1:
        pil_images[0].save(pdf_buffer, format="PDF")
    else:
        pil_images[0].save(
            pdf_buffer,
            format="PDF",
            save_all=True,
            append_images=pil_images[1:],
        )

    pdf_buffer.seek(0)
    return pdf_buffer


def images_to_docx(image_files):
    doc = Document()
    doc.add_heading("EduSolve AI", 0)
    doc.add_paragraph(f"Total Images: {len(image_files)}")

    for idx, img_file in enumerate(image_files):
        try:
            img_file.seek(0)
            image = Image.open(img_file)
        except Exception:
            img_file.seek(0)
            image = Image.open(io.BytesIO(img_file.read()))

        temp_buffer = io.BytesIO()
        if image.mode in ("RGBA", "P"):
            image = image.convert("RGB")
        image.save(temp_buffer, format="PNG")
        temp_buffer.seek(0)

        doc.add_heading(f"Image {idx + 1}", level=2)
        doc.add_picture(temp_buffer, width=Inches(5.5))

    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    doc_buffer.seek(0)
    return doc_buffer

# ============================================================
# NEXT-LEVEL SOURCE STUDIO FEATURES
# ============================================================

def _source_context(query, k=8, per_doc=1800):
    """Return grounded context from the currently indexed knowledge base.

    This helper is intentionally safe: if no files are indexed, it returns a
    clear marker instead of allowing the new Source Studio features to crash.
    """
    try:
        docs = hybrid_search(query, k=k)
    except Exception as e:
        print(f"Source retrieval error: {e}")
        return "No indexed source context is available."

    if not docs:
        return "No indexed source context is available."

    blocks = []
    for i, doc in enumerate(docs, 1):
        meta = getattr(doc, "metadata", {}) or {}
        if not isinstance(meta, dict):
            meta = {}
        src = meta.get("source", "Unknown source")
        page = meta.get("page", meta.get("page_number", ""))
        label = f"Source {i}: {os.path.basename(str(src))}"
        if page not in ("", None):
            label += f" | page {page}"
        text = _clean_text(getattr(doc, "page_content", ""))
        if text:
            blocks.append(f"[{label}]\n{text[:per_doc]}")
    return "\n\n".join(blocks) if blocks else "No indexed source context is available."


def _has_source_material():
    """Fast check used by Source Studio generators."""
    return st.session_state.get("vectordb") is not None


def get_source_inventory():
    """Return a compact inventory of indexed documents/chunks without crashing the UI."""
    db = st.session_state.get("vectordb")
    if db is None:
        return {"total_chunks": 0, "sources": []}
    try:
        docs = list(db.docstore._dict.values())
    except Exception:
        docs = []

    # FAISS/docstore internals can vary between langchain-community versions.
    # If the docstore cannot be read, use the metadata already stored by process_files.
    if not docs:
        names = st.session_state.get("kb_files", []) or []
        chunk_count = int(st.session_state.get("kb_chunk_count", 0) or 0)
        if names:
            each = max(1, chunk_count // len(names)) if chunk_count else 0
            return {
                "total_chunks": chunk_count,
                "sources": [{"name": os.path.basename(str(n)), "chunks": each, "pages": []} for n in names],
            }

    grouped = {}
    for d in docs:
        meta = getattr(d, "metadata", {}) or {}
        src = str(meta.get("source", "Unknown source"))
        name = os.path.basename(src) or "Unknown source"
        grouped.setdefault(name, {"name": name, "chunks": 0, "pages": set()})
        grouped[name]["chunks"] += 1
        p = meta.get("page", meta.get("page_number"))
        if p is not None:
            grouped[name]["pages"].add(str(p))

    sources = []
    for item in grouped.values():
        item["pages"] = sorted(
            item["pages"], key=lambda x: int(x) if str(x).isdigit() else str(x)
        )
        sources.append(item)
    return {
        "total_chunks": len(docs),
        "sources": sorted(sources, key=lambda x: x["name"].lower()),
    }


def _studio_json(prompt, fallback):
    try:
        data = _json_from_text(_invoke(prompt, temperature=0.35), default=None)
        return data if data is not None else fallback
    except Exception as e:
        print(f"Source Studio LLM error: {e}")
        return fallback


def generate_source_brief(language="Hinglish"):
    inv = get_source_inventory()
    if not inv.get("sources"):
        return "## 📚 Source Brief\n\nPehle Knowledge Base me PDF/TXT/CSV upload karo."

    context = _source_context(
        "main topics definitions concepts formulas examples important facts",
        k=10,
    )
    prompt = f"""You are a source-grounded academic assistant. Use ONLY the supplied source context.
Create a concise but useful source brief in {language}.
Include: overview, major topics, key concepts, formulas/facts if present, likely exam areas, and source coverage.
Do not invent facts. If something is not present, say it is not found.
Return plain text with clear headings.
Inventory: {json.dumps(inv, ensure_ascii=False)}
SOURCE CONTEXT:\n{context}"""
    try:
        answer = _invoke(prompt, temperature=0.25)
        if answer and str(answer).strip():
            return str(answer).strip()
    except Exception as e:
        print(f"Source Brief LLM error: {e}")

    lines = ["## 📚 Source Brief", "", f"**Indexed chunks:** {inv.get('total_chunks', 0)}", "", "### Sources"]
    for item in inv.get("sources", []):
        pages = item.get("pages") or []
        page_text = f" | pages: {', '.join(map(str, pages[:12]))}" if pages else ""
        lines.append(f"- **{item.get('name', 'Unknown source')}** — {item.get('chunks', 0)} chunks{page_text}")
    lines += ["", "### Source Extract", context[:7000] if context else "No source text could be retrieved."]
    lines += ["", "> AI synthesis was temporarily unavailable; this brief is generated directly from the indexed source metadata/context."]
    return "\n".join(lines)


def generate_flashcards(topic="all material", num_cards=10, language="Hinglish"):
    n = max(3, min(30, int(num_cards)))
    if not _has_source_material():
        return []

    context = _source_context(topic, k=10)
    fallback = [
        {
            "front": f"{topic} — Card {i + 1}",
            "back": "Review the uploaded source material for this concept.",
            "difficulty": "Medium",
            "source": "Uploaded material",
        }
        for i in range(n)
    ]
    prompt = f"""Create {n} high-quality study flashcards from ONLY the source context.
Language: {language}. Topic: {topic}.
Return ONLY JSON array. Each object MUST contain:
{{"front":"question/term","back":"answer","difficulty":"Easy|Medium|Hard","source":"source filename or Uploaded material"}}
Avoid duplicates. Keep answers concise and source-grounded.
SOURCE CONTEXT:\n{context}"""
    data = _studio_json(prompt, fallback)
    if not isinstance(data, list):
        return fallback

    normalized = []
    for i, card in enumerate(data[:n]):
        if not isinstance(card, dict):
            continue
        normalized.append(
            {
                "front": str(card.get("front") or card.get("question") or f"{topic} — Card {i + 1}"),
                "back": str(card.get("back") or card.get("answer") or "Review the uploaded source material."),
                "difficulty": str(card.get("difficulty") or "Medium"),
                "source": str(card.get("source") or "Uploaded material"),
            }
        )
    return normalized or fallback


def generate_mind_map(topic, language="Hinglish"):
    topic = _clean_text(topic) or "General Topic"
    has_source = _has_source_material()
    if has_source:
        context = _source_context(topic, k=12)
        mode_instruction = "Use the supplied source context as the primary factual basis. Do not add source-specific facts that are not supported."
    else:
        context = "No uploaded source is available. Use your general academic knowledge and clearly keep the map conceptual."
        mode_instruction = "No source is available, so create a general educational concept map."

    fallback = {
        "topic": topic,
        "branches": [
            {"name": "Definition / Meaning", "children": [f"Understand what {topic} means"]},
            {"name": "Core Concepts", "children": ["Main ideas", "Key terms", "Important relationships"]},
            {"name": "Working / Process", "children": ["Step-by-step flow", "Inputs and outputs"]},
            {"name": "Examples / Applications", "children": ["Practical example", "Real-world use"]},
            {"name": "Exam & Revision", "children": ["Important points", "Common mistakes"]},
        ],
        "message": "General concept map generated." if not has_source else "Source-grounded concept map generated.",
    }

    prompt = f"""Create a high-quality educational mind map for: {topic}
Language: {language}
{mode_instruction}

Return ONLY valid JSON with exactly this shape:
{{
  "topic": "{topic}",
  "branches": [
    {{"name": "branch name", "children": ["child 1", "child 2", "child 3"]}}
  ]
}}

Requirements:
- 5 to 8 meaningful branches.
- 2 to 6 short children per branch.
- Cover definition, core concepts, working/process, examples/applications, comparison/relationships when relevant, mistakes and revision/exam points.
- No markdown, no code fences, no commentary outside JSON.
- Keep each child concise but informative.

SOURCE/KNOWLEDGE CONTEXT:
{context[:MAX_CONTEXT_CHARS]}"""
    data = _studio_json(prompt, fallback)
    if not isinstance(data, dict):
        return fallback

    branches = data.get("branches", [])
    if isinstance(branches, dict):
        branches = [branches]
    normalized = []
    if isinstance(branches, list):
        for branch in branches:
            if isinstance(branch, str):
                normalized.append({"name": branch, "children": []})
                continue
            if not isinstance(branch, dict):
                continue
            children = branch.get("children", branch.get("items", []))
            if isinstance(children, str):
                children = [children]
            if not isinstance(children, list):
                children = []
            clean_children = []
            for child in children:
                if isinstance(child, dict):
                    child = child.get("name") or child.get("text") or child.get("point") or ""
                text = _clean_text(child)
                if text:
                    clean_children.append(text)
            normalized.append({
                "name": _clean_text(branch.get("name") or branch.get("title") or "Key Concept") or "Key Concept",
                "children": clean_children[:8],
            })
    normalized = [b for b in normalized if b.get("name")]
    if not normalized:
        return fallback
    return {
        "topic": _clean_text(data.get("topic") or topic) or topic,
        "branches": normalized[:10],
        "message": "Source-grounded concept map generated." if has_source else "General AI concept map generated — no source required.",
    }


def generate_exam_paper(topic="all material", num_questions=10, difficulty="Mixed", language="Hinglish"):
    n = max(3, min(30, int(num_questions)))
    if not _has_source_material():
        return []

    context = _source_context(topic, k=12)
    fallback = [
        {
            "question": f"Explain an important concept from {topic}.",
            "answer": "Refer to the uploaded source material.",
            "difficulty": difficulty,
            "type": "Short Answer",
            "marks": 5,
            "topic": topic,
        }
        for _ in range(n)
    ]
    prompt = f"""Generate an exam paper from ONLY the source material.
Topic: {topic}; Questions: {n}; Difficulty: {difficulty}; Language: {language}.
Return ONLY JSON array. Every object MUST contain:
{{"question":"...","answer":"...","difficulty":"Easy|Medium|Hard|Mixed","type":"MCQ|Short Answer|Long Answer","marks":2|5|10,"topic":"..."}}
Cover different concepts, avoid duplicates, and never invent source-specific facts.
SOURCE CONTEXT:\n{context}"""
    data = _studio_json(prompt, fallback)
    if not isinstance(data, list):
        return fallback

    normalized = []
    for i, q in enumerate(data[:n]):
        if not isinstance(q, dict):
            continue
        try:
            marks = int(q.get("marks", 5))
        except Exception:
            marks = 5
        normalized.append(
            {
                "question": str(q.get("question") or f"Question {i + 1} from {topic}"),
                "answer": str(q.get("answer") or "Refer to the uploaded source material."),
                "difficulty": str(q.get("difficulty") or difficulty),
                "type": str(q.get("type") or "Short Answer"),
                "marks": max(1, marks),
                "topic": str(q.get("topic") or topic),
            }
        )
    return normalized or fallback


def compare_source_topics(topic_a, topic_b, language="Hinglish"):
    if not _has_source_material():
        return {
            "topic_a": topic_a,
            "topic_b": topic_b,
            "similarities": [],
            "differences": [],
            "summary": "Please upload study material first.",
        }

    context = _source_context(f"{topic_a} and {topic_b}", k=12)
    fallback = {
        "topic_a": topic_a,
        "topic_b": topic_b,
        "similarities": [],
        "differences": [],
        "summary": "Compare the two topics using the uploaded material.",
    }
    prompt = f"""Compare {topic_a} and {topic_b} using ONLY the source context.
Language: {language}.
Return ONLY JSON with keys topic_a, topic_b, similarities (array), differences (array of objects with aspect,a,b), summary.
If the source does not support a comparison, state that clearly.
SOURCE CONTEXT:\n{context}"""
    data = _studio_json(prompt, fallback)
    if not isinstance(data, dict):
        return fallback

    similarities = data.get("similarities", [])
    differences = data.get("differences", [])
    if isinstance(similarities, str):
        similarities = [similarities]
    if not isinstance(similarities, list):
        similarities = []
    if not isinstance(differences, list):
        differences = []

    norm_diff = []
    for d in differences:
        if not isinstance(d, dict):
            continue
        norm_diff.append(
            {
                "aspect": str(d.get("aspect") or "Aspect"),
                "a": str(d.get("a") or "Not found in source"),
                "b": str(d.get("b") or "Not found in source"),
            }
        )
    return {
        "topic_a": str(data.get("topic_a") or topic_a),
        "topic_b": str(data.get("topic_b") or topic_b),
        "similarities": [str(x) for x in similarities],
        "differences": norm_diff,
        "summary": str(data.get("summary") or "No supported comparison summary found."),
    }


def generate_study_guide(topic="all material", language="Hinglish"):
    topic = (str(topic or "all material").strip() or "all material")
    if not _has_source_material():
        return "## 📘 Study Guide\n\nPehle Knowledge Base me PDF/TXT/CSV upload karo."

    context = _source_context(topic, k=12)
    if not context or context.startswith("No indexed source context"):
        return f"## 📘 Study Guide\n\n**{topic}** ke liye source me relevant material nahi mila."

    prompt = f"""Create a practical study guide for {topic} in {language}, grounded ONLY in the source context.
Include: prerequisites, learning objectives, concepts in order, examples/applications found in source, common mistakes, revision checklist, and exam focus.
Use headings and bullets. Do not invent unsupported information.
SOURCE CONTEXT:\n{context}"""
    try:
        answer = _invoke(prompt, temperature=0.25)
        if answer and str(answer).strip():
            return str(answer).strip()
    except Exception as e:
        print(f"Study Guide LLM error: {e}")

    return f"""## 📘 Study Guide — {topic}

### Learning objectives
- Identify the definitions and core concepts in the uploaded material.
- Practice the examples, comparisons and processes actually present in the source.

### Source-based concepts
{context[:7000]}

### Revision checklist
- [ ] Definitions
- [ ] Core concepts
- [ ] Examples/applications
- [ ] Comparisons
- [ ] Diagrams/processes
- [ ] Self-test questions

> AI synthesis was temporarily unavailable; this guide is generated directly from the retrieved source context.
"""


def build_study_pack(topic="all material", language="Hinglish"):
    brief = generate_source_brief(language)
    guide = generate_study_guide(topic, language)
    cards = generate_flashcards(topic, 12, language)
    doc = Document()
    doc.add_heading("EduSolve AI — Study Pack", 0)
    doc.add_paragraph(f"Topic: {topic}")
    doc.add_heading("Source Brief", level=1)
    doc.add_paragraph(str(brief))
    doc.add_heading("Study Guide", level=1)
    doc.add_paragraph(str(guide))
    doc.add_heading("Flashcards", level=1)
    for i, card in enumerate(cards, 1):
        doc.add_paragraph(f"{i}. {card.get('front','')}", style="List Number")
        doc.add_paragraph(str(card.get('back','')))
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
